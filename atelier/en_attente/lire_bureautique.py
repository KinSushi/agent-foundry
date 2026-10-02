"""Une révision Word non acceptée est invisible à qui lit le texte « normalement » :
sur un paragraphe portant une insertion suivie (w:ins « lundi ») et une
suppression suivie (w:del « vendredi »), python-docx 1.2.0 rend
`doc.paragraphs[1].text` = 'Le prestataire livre le avant midi.' (35 caractères,
mesuré) — ni la version acceptée, ni l'originale, et aucun signe qu'une révision
attend. L'interpréteur propre n'a ni python-docx ni python-pptx (mesuré :
`python3.14 -c 'import docx'` rend ModuleNotFoundError).

QUESTION
    Que dit ce document Word, PowerPoint ou OpenDocument, et contient-il des
    révisions non acceptées ou des commentaires ?
MESURE
    Lecture des parties XML de l'archive (zipfile + xml.etree), sans rien exécuter.
    .docx/.docm/.dotx : paragraphes du corps (texte avec révisions acceptées, et
    texte original quand il diffère), titres par style (niveau de plan du style ou
    de sa lignée, sinon nom « Heading N »/« Titre N »), tableaux, révisions suivies
    (w:ins, w:del, w:moveFrom/moveTo, changements de mise en forme *PrChange) avec
    auteur et date, dans le corps, les en-têtes, pieds de page et notes ;
    commentaires (comments.xml, état résolu d'après commentsExtended.xml) ; suivi
    des modifications activé ou non (settings.xml). .pptx : texte par diapositive,
    titre, notes, diapositives masquées, commentaires (anciens et modernes).
    .odt/.odp : titres (text:h), paragraphes, tableaux, text:tracked-changes,
    office:annotation (résolue si loext:resolved). Moteurs python-docx et
    python-pptx s'ils sont installés (noms de styles, commentaires, formes).
HYPOTHÈSES
    Le document est un paquet Office Open XML ou OpenDocument non chiffré. Une
    révision présente dans le XML est non acceptée (Word retire le balisage d'une
    révision acceptée). Un commentaire présent compte, même résolu.
LIMITES
    Ne lit ni .doc ni .ppt (binaires), ni les documents chiffrés. Ne restitue pas
    la mise en page, les images, les champs calculés (seul leur dernier résultat
    affiché), ni le texte des zones de dessin d'un .pptx hors forme texte. En
    OpenDocument, le texte original d'une suppression n'est pas réinséré dans le
    paragraphe : il est donné dans la liste des révisions. Chaque partie XML est
    chargée entière en mémoire, plafonnée par --max-octets.
CONTRE-EXEMPLES
    idxexample.odt (livré avec LibreOffice 24.2) contient 34 éléments text:p et
    text:h (comptés avec xml.etree) ; l'outil annonce 33 paragraphes, car le
    paragraphe d'une zone de texte (draw:text-box) ancrée dans un paragraphe est
    fondu dans celui-ci. Un commentaire n'est pas rattaché au passage commenté :
    seul le numéro du paragraphe qui porte sa référence est donné.
INVOCATION
    {outil} --xml '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Livraison le </w:t></w:r><w:ins w:id="1" w:author="Alice" w:date="2026-09-30T10:00:00Z"><w:r><w:t>lundi</w:t></w:r></w:ins></w:p></w:body></w:document>' --json
DOMAINE
    Documents .docx/.docm/.dotx, .pptx/.pptm/.potx, .odt/.ott, .odp/.otp à relire
    ou à transmettre : savoir ce qu'ils disent et s'ils portent encore des
    révisions ou des commentaires avant envoi, signature ou extraction.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
import zlib
import xml.etree.ElementTree as ET
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from types import MappingProxyType
from typing import Any, Iterator

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

try:
    import docx  # type: ignore[import-not-found]
    from docx.text.paragraph import Paragraph as ParagrapheDocx  # type: ignore[import-not-found]
except ImportError:
    docx = None
    ParagrapheDocx = None

try:
    import pptx  # type: ignore[import-not-found]
except ImportError:
    pptx = None

RACINE = Path(__file__).resolve().parent

INTITULE_QUESTION = "QUESTION"
INTITULE_MESURE = "MESURE"
INTITULE_HYPOTHESES = "HYPOTHÈSES"
INTITULE_LIMITES = "LIMITES"
INTITULE_CONTRE_EXEMPLES = "CONTRE-EXEMPLES"
INTITULE_INVOCATION = "INVOCATION"
INTITULE_DOMAINE = "DOMAINE"
INTITULES = (
    INTITULE_QUESTION, INTITULE_MESURE, INTITULE_HYPOTHESES, INTITULE_LIMITES,
    INTITULE_CONTRE_EXEMPLES, INTITULE_INVOCATION, INTITULE_DOMAINE,
)

EXTENSIONS_WORD = (".docx", ".docm", ".dotx", ".dotm")
EXTENSIONS_POWERPOINT = (".pptx", ".pptm", ".potx", ".potm")
EXTENSIONS_ODF = (".odt", ".ott", ".odp", ".otp")
EXTENSIONS_ODF_PLAT = (".fodt", ".fodp")
EXTENSIONS_BINAIRES = (".doc", ".ppt", ".dot", ".pps")
NS_OFFICE = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
NS_TEXT = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
NS_TABLE = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
NS_DRAW = "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
NS_PRESENTATION = "urn:oasis:names:tc:opendocument:xmlns:presentation:1.0"
NS_STYLE = "urn:oasis:names:tc:opendocument:xmlns:style:1.0"
NS_DC = "http://purl.org/dc/elements/1.1/"
SIGNATURE_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
REVISIONS_WORD = MappingProxyType({
    "ins": "insertion", "cellIns": "insertion", "moveTo": "deplacement",
    "del": "suppression", "cellDel": "suppression", "moveFrom": "deplacement",
    "rPrChange": "mise_en_forme", "pPrChange": "mise_en_forme",
    "sectPrChange": "mise_en_forme", "tblPrChange": "mise_en_forme",
    "tblPrExChange": "mise_en_forme", "tblGridChange": "mise_en_forme",
    "trPrChange": "mise_en_forme", "tcPrChange": "mise_en_forme",
    "numberingChange": "mise_en_forme", "cellMerge": "mise_en_forme",
})
IGNORES_WORD = frozenset({
    "pPr", "rPr", "instrText", "delInstrText", "txbxContent", "Fallback",
    "commentReference", "annotationRef", "footnoteReference", "endnoteReference",
})
MOTIF_TITRE = re.compile(r"(heading|titre|title|überschrift|título|titolo)\s*(\d)", re.IGNORECASE)
LIMITE_EXAMINES = 200
LIMITE_EXTRAIT = 160


class ErreurDocument(Exception):
    """Entrée illisible ou refusée ; `code` est le code de sortie proposé."""

    def __init__(self, message: str, code: int = 2) -> None:
        super().__init__(message)
        self.code = code


ERREURS_LECTURE = (
    OSError, zipfile.BadZipFile, zipfile.LargeZipFile, ET.ParseError, KeyError,
    ValueError, EOFError, zlib.error, ErreurDocument,
)


@dataclass
class Reglages:
    """Paramètres communs à toutes les lectures."""

    moteur: str
    max_octets: int
    max_paragraphes: int


@dataclass
class Releve:
    """Ce qu'on a lu d'un document, quel que soit son format."""

    format: str
    moteur: str = "stdlib"
    paragraphes: list[dict[str, Any]] = field(default_factory=list)
    titres: list[dict[str, Any]] = field(default_factory=list)
    tableaux: list[dict[str, Any]] = field(default_factory=list)
    revisions: list[dict[str, Any]] = field(default_factory=list)
    commentaires: list[dict[str, Any]] = field(default_factory=list)
    diapositives: list[dict[str, Any]] = field(default_factory=list)
    parties: list[str] = field(default_factory=list)
    suivi_active: bool | None = None
    notes: list[str] = field(default_factory=list)

    def ajouter_paragraphe(self, texte: str, **details: Any) -> None:
        """Retient un paragraphe et, s'il en est un, le titre qu'il porte."""
        self.paragraphes.append({"n": len(self.paragraphes) + 1, "texte": texte, **details})
        if details.get("niveau") is not None and texte.strip():
            self.titres.append({"niveau": details["niveau"], "texte": texte.strip()})


# --------------------------------------------------------------------------- #
# Outils de base

def extraire_contrat(doc: str) -> dict[str, str]:
    """Découpe la docstring en sections selon les intitulés du contrat."""
    contrat: dict[str, str] = {}
    courant: str | None = None
    for ligne in doc.splitlines():
        tete = ligne.strip()
        if tete in INTITULES:
            courant = tete
            contrat[courant] = ""
        elif courant is not None:
            contrat[courant] = (contrat[courant] + " " + tete).strip()
    return contrat


def nom_local(etiquette: Any) -> str:
    """Nom local d'une étiquette ; chaîne vide pour un commentaire XML lxml."""
    if not isinstance(etiquette, str):
        return ""
    return etiquette.rsplit("}", 1)[-1]


def attribut_relation(element: Any) -> str | None:
    """Attribut r:id (espace de noms des relations, transitionnel ou strict)."""
    for cle, valeur in element.attrib.items():
        if cle.endswith("}id") and "relationships" in cle:
            return valeur
    return None


def attribut_local(element: Any, nom: str) -> str | None:
    """Attribut lu par son nom local, quel que soit son espace de noms."""
    for cle, valeur in element.attrib.items():
        if nom_local(cle) == nom:
            return valeur
    return None


def extrait(texte: str) -> str:
    """Texte borné pour les listes de révisions et de commentaires."""
    texte = " ".join(texte.split())
    return texte if len(texte) <= LIMITE_EXTRAIT else texte[: LIMITE_EXTRAIT - 1] + "…"


def verifier_dtd(donnees: bytes, nom: str) -> None:
    """Refuse toute déclaration DTD : jamais légitime dans ces formats."""
    if b"<!DOCTYPE" in donnees or b"<!ENTITY" in donnees:
        raise ErreurDocument(f"{nom} : déclaration DTD refusée (jamais légitime ici)", 1)


def analyser_xml(donnees: bytes, nom: str) -> ET.Element:
    """Analyse un XML déjà lu, après refus des DTD."""
    verifier_dtd(donnees, nom)
    try:
        return ET.fromstring(donnees)
    except ET.ParseError as exc:
        raise ErreurDocument(f"{nom} : XML mal formé ({exc})") from exc


def lire_partie(archive: zipfile.ZipFile, nom: str, plafond: int) -> ET.Element:
    """Lit une partie XML bornée en taille décompressée (déclarée puis réelle)."""
    info = archive.getinfo(nom)
    if info.file_size > plafond:
        raise ErreurDocument(
            f"{nom} déclare {info.file_size} octets décompressés, au-delà du plafond {plafond} "
            "(garde contre les bombes zip ; voir --max-octets)", 1)
    with archive.open(info) as flux:
        donnees = flux.read(plafond + 1)
    if len(donnees) > plafond:
        raise ErreurDocument(f"{nom} dépasse {plafond} octets décompressés (voir --max-octets)", 1)
    return analyser_xml(donnees, nom)


def ouvrir_archive(chemin: Path) -> zipfile.ZipFile:
    """Ouvre le paquet ou explique pourquoi ce n'en est pas un."""
    with chemin.open("rb") as brut:
        signature = brut.read(8)
    if signature == SIGNATURE_OLE:
        raise ErreurDocument(f"{chemin.name} est un conteneur OLE (document chiffré ou format binaire .doc/.ppt)")
    try:
        return zipfile.ZipFile(chemin)
    except zipfile.BadZipFile as exc:
        raise ErreurDocument(f"{chemin.name} n'est pas une archive zip valide ({exc})") from exc


def relations(archive: zipfile.ZipFile, partie: str, plafond: int) -> list[tuple[str, str, str]]:
    """(identifiant, type abrégé, cible résolue) des relations internes d'une partie."""
    dossier, _, nom = partie.rpartition("/")
    chemin_rels = f"{dossier}/_rels/{nom}.rels" if dossier else f"_rels/{nom}.rels"
    if chemin_rels not in archive.namelist():
        return []
    resultat = []
    for rel in lire_partie(archive, chemin_rels, plafond):
        if rel.get("TargetMode") == "External":
            continue
        cible = rel.get("Target", "")
        resolue = cible.lstrip("/") if cible.startswith("/") else normaliser_chemin(dossier, cible)
        resultat.append((rel.get("Id", ""), rel.get("Type", "").rsplit("/", 1)[-1], resolue))
    return resultat


def normaliser_chemin(dossier: str, cible: str) -> str:
    """Résout ../ dans un chemin interne à l'archive."""
    morceaux = [m for m in dossier.split("/") if m]
    for morceau in cible.split("/"):
        if morceau == "..":
            if morceaux:
                morceaux.pop()
        elif morceau not in ("", "."):
            morceaux.append(morceau)
    return "/".join(morceaux)


def partie_principale(archive: zipfile.ZipFile, plafond: int, defaut: str) -> str:
    """Partie officeDocument désignée par _rels/.rels."""
    for _ident, type_rel, cible in relations(archive, "", plafond):
        if type_rel == "officeDocument":
            return cible
    return defaut


# --------------------------------------------------------------------------- #
# WordprocessingML : texte, styles, révisions

def textes_paragraphe(paragraphe: Any) -> tuple[str, str]:
    """(texte révisions acceptées, texte original) d'un paragraphe Word."""
    acceptes: list[str] = []
    originaux: list[str] = []
    parcourir_runs(paragraphe, False, False, acceptes, originaux)
    return "".join(acceptes), "".join(originaux)


def parcourir_runs(element: Any, insere: bool, supprime: bool, acceptes: list[str], originaux: list[str]) -> None:
    """Descend dans les runs ; insertions et suppressions vont chacune de leur côté."""
    for enfant in element:
        local = nom_local(enfant.tag)
        if local in IGNORES_WORD or not local:
            continue
        morceau = morceau_run(local, enfant)
        if morceau is not None:
            if not supprime and local != "delText":
                acceptes.append(morceau)
            if not insere:
                originaux.append(morceau)
        elif local in ("ins", "moveTo"):
            parcourir_runs(enfant, True, supprime, acceptes, originaux)
        elif local in ("del", "moveFrom"):
            parcourir_runs(enfant, insere, True, acceptes, originaux)
        elif local == "AlternateContent":
            parcourir_runs(enfant[:1] if len(enfant) else [], insere, supprime, acceptes, originaux)
        else:
            parcourir_runs(enfant, insere, supprime, acceptes, originaux)


def morceau_run(local: str, element: Any) -> str | None:
    """Texte produit par un élément feuille d'un run ; None si ce n'en est pas un."""
    if local in ("t", "delText"):
        return element.text or ""
    if local == "tab":
        return "\t"
    if local in ("br", "cr"):
        return "\n"
    if local == "noBreakHyphen":
        return "-"
    return None


@dataclass
class Styles:
    """Styles de paragraphe : identifiant -> (nom, niveau de titre)."""

    noms: dict[str, str] = field(default_factory=dict)
    niveaux: dict[str, int | None] = field(default_factory=dict)
    defaut: str | None = None

    def nom(self, ident: str | None) -> str | None:
        """Nom lisible du style (style par défaut si aucun n'est donné)."""
        ident = ident or self.defaut
        if ident is None:
            return None
        return self.noms.get(ident, ident)

    def niveau(self, ident: str | None) -> int | None:
        """Niveau de titre du style, ou d'après son nom ou son identifiant."""
        ident = ident or self.defaut
        if ident is None:
            return None
        if ident in self.niveaux:
            return self.niveaux[ident]
        return niveau_par_nom(ident)


def niveau_par_nom(nom: str) -> int | None:
    """« Heading 2 », « Titre2 » -> 2 ; « Title » -> 0 ; sinon None."""
    correspondance = MOTIF_TITRE.fullmatch(nom.strip())
    if correspondance:
        return int(correspondance.group(2))
    return 0 if nom.strip().lower() in ("title", "titre") else None


def lire_styles(racine: ET.Element | None) -> Styles:
    """Construit la table des styles de paragraphe, lignées basedOn comprises."""
    styles = Styles()
    if racine is None:
        return styles
    bruts: dict[str, dict[str, Any]] = {}
    for style in racine:
        if nom_local(style.tag) != "style" or attribut_local(style, "type") != "paragraph":
            continue
        ident = attribut_local(style, "styleId") or ""
        bruts[ident] = description_style(style)
        if attribut_local(style, "default") in ("1", "true"):
            styles.defaut = ident
    for ident, brut in bruts.items():
        styles.noms[ident] = nom_affiche(brut["nom"] or ident)
        styles.niveaux[ident] = niveau_herite(ident, bruts)
    return styles


def nom_affiche(nom: str) -> str:
    """Nom montré par Word pour les styles intégrés écrits en minuscules (heading 1 -> Heading 1)."""
    if nom in ("caption", "footer", "header") or re.fullmatch(r"heading [1-9]", nom):
        return nom[0].upper() + nom[1:]
    return nom


def description_style(style: ET.Element) -> dict[str, Any]:
    """Nom, parent et niveau de plan déclarés par un w:style."""
    description: dict[str, Any] = {"nom": None, "parent": None, "plan": None}
    for enfant in style.iter():
        local = nom_local(enfant.tag)
        if local == "name":
            description["nom"] = attribut_local(enfant, "val")
        elif local == "basedOn":
            description["parent"] = attribut_local(enfant, "val")
        elif local == "outlineLvl":
            description["plan"] = int(attribut_local(enfant, "val") or "9")
    return description


def niveau_herite(ident: str, bruts: dict[str, dict[str, Any]]) -> int | None:
    """Niveau de titre d'un style : plan déclaré, sinon nom, en remontant la lignée."""
    vus: set[str] = set()
    courant: str | None = ident
    while courant in bruts and courant not in vus:
        vus.add(courant)
        brut = bruts[courant]
        if brut["plan"] is not None:
            return brut["plan"] + 1 if brut["plan"] < 9 else None
        niveau = niveau_par_nom(brut["nom"] or courant)
        if niveau is not None:
            return niveau
        courant = brut["parent"]
    return None


def style_paragraphe(paragraphe: Any) -> tuple[str | None, int | None]:
    """(identifiant de style, niveau de plan direct) lus dans w:pPr."""
    ident, plan = None, None
    for enfant in paragraphe:
        if nom_local(enfant.tag) != "pPr":
            continue
        for propriete in enfant:
            local = nom_local(propriete.tag)
            if local == "pStyle":
                ident = attribut_local(propriete, "val")
            elif local == "outlineLvl":
                plan = int(attribut_local(propriete, "val") or "9")
    return ident, (plan + 1 if plan is not None and plan < 9 else None)


def revisions_word(racine: Any, partie: str, ancrage: dict[int, int]) -> list[dict[str, Any]]:
    """Toutes les révisions suivies d'une partie, dans l'ordre du document."""
    trouvees = []
    for element in racine.iter():
        local = nom_local(element.tag)
        if local not in REVISIONS_WORD:
            continue
        trouvees.append({
            "type": REVISIONS_WORD[local],
            "balise": f"w:{local}",
            "auteur": attribut_local(element, "author"),
            "date": attribut_local(element, "date"),
            "texte": extrait(texte_revision(element, local)),
            "partie": partie,
            "paragraphe": ancrage.get(id(element)),
        })
    return trouvees


def texte_revision(element: Any, local: str) -> str:
    """Texte inséré ou supprimé ; « ¶ » pour une marque de paragraphe."""
    if local.endswith("Change") or local == "cellMerge":
        return ""
    morceaux = [e.text or "" for e in element.iter() if nom_local(e.tag) in ("t", "delText")]
    return "".join(morceaux) or "¶"


# --------------------------------------------------------------------------- #
# Word : corps, commentaires, parties annexes

def lire_corps_word(racine: Any, styles: Styles, releve: Releve, partie: str) -> None:
    """Parcourt le corps (paragraphes, tableaux, zones de texte) et ses révisions."""
    ancrage: dict[int, int] = {}
    references: dict[str, int] = {}
    corps = next((e for e in racine if nom_local(e.tag) == "body"), racine)
    for bloc in blocs_word(corps):
        if nom_local(bloc.tag) == "p":
            traiter_paragraphe_word(bloc, styles, releve, ancrage, references)
        else:
            traiter_tableau_word(bloc, styles, releve, ancrage, references)
    for zone in zones_de_texte(racine):
        for paragraphe in (p for p in zone.iter() if nom_local(p.tag) == "p"):
            traiter_paragraphe_word(paragraphe, styles, releve, ancrage, references, zone_texte=True)
    releve.revisions.extend(revisions_word(racine, partie, ancrage))
    if not releve.commentaires:
        releve.commentaires = [
            {"id": i, "auteur": None, "date": None, "texte": "(texte dans comments.xml, absent)", "resolu": None}
            for i in references
        ]
    for commentaire in releve.commentaires:
        commentaire["paragraphe"] = references.get(commentaire["id"])


def blocs_word(conteneur: Any) -> Iterator[Any]:
    """Paragraphes et tableaux de premier niveau, contrôles de contenu dépliés."""
    for enfant in conteneur:
        local = nom_local(enfant.tag)
        if local in ("p", "tbl"):
            yield enfant
        elif local in ("sdt", "sdtContent", "customXml", "ins", "del", "moveTo", "moveFrom"):
            yield from blocs_word(enfant)


def zones_de_texte(racine: Any) -> list[Any]:
    """Zones de texte (w:txbxContent) hors variantes de repli VML."""
    exclues = {id(e) for repli in racine.iter() if nom_local(repli.tag) == "Fallback" for e in repli.iter()}
    return [e for e in racine.iter() if nom_local(e.tag) == "txbxContent" and id(e) not in exclues]


def traiter_paragraphe_word(
    paragraphe: Any, styles: Styles, releve: Releve, ancrage: dict[int, int],
    references: dict[str, int], tableau: int | None = None, zone_texte: bool = False,
) -> None:
    """Retient le texte d'un paragraphe et ancre ses révisions et commentaires."""
    accepte, original = textes_paragraphe(paragraphe)
    ident, plan = style_paragraphe(paragraphe)
    numero = len(releve.paragraphes) + 1
    for element in paragraphe.iter():
        local = nom_local(element.tag)
        if local in REVISIONS_WORD:
            ancrage[id(element)] = numero
        elif local == "commentReference":
            references.setdefault(attribut_local(element, "id") or "?", numero)
    details: dict[str, Any] = {"style": styles.nom(ident), "niveau": plan if plan is not None else styles.niveau(ident)}
    if original != accepte:
        details["original"] = original
    if tableau is not None:
        details["tableau"] = tableau
    if zone_texte:
        details["zone_de_texte"] = True
        details["niveau"] = None
    releve.ajouter_paragraphe(accepte, **details)


def traiter_tableau_word(
    tableau: Any, styles: Styles, releve: Releve, ancrage: dict[int, int], references: dict[str, int],
) -> None:
    """Dimensions du tableau, première ligne, et paragraphes de ses cellules."""
    index = len(releve.tableaux) + 1
    fiche: dict[str, Any] = {"index": index}
    releve.tableaux.append(fiche)
    lignes = [e for e in tableau if nom_local(e.tag) == "tr"]
    largeurs = [len([c for c in ligne if nom_local(c.tag) == "tc"]) for ligne in lignes]
    premiere = []
    for numero, ligne in enumerate(lignes):
        for cellule in (c for c in ligne if nom_local(c.tag) == "tc"):
            textes = []
            for bloc in blocs_word(cellule):
                if nom_local(bloc.tag) == "p":
                    traiter_paragraphe_word(bloc, styles, releve, ancrage, references, tableau=index)
                    textes.append(releve.paragraphes[-1]["texte"])
                else:
                    traiter_tableau_word(bloc, styles, releve, ancrage, references)
            if numero == 0:
                premiere.append(extrait(" ".join(textes)))
    fiche.update(lignes=len(lignes), colonnes=max(largeurs, default=0), premiere_ligne=premiere)


def lire_commentaires_word(archive: zipfile.ZipFile, cibles: dict[str, str], plafond: int) -> list[dict[str, Any]]:
    """Commentaires de comments.xml, avec leur état résolu si commentsExtended existe."""
    nom = cibles.get("comments")
    if nom is None or nom not in archive.namelist():
        return []
    resolus = etats_resolus(archive, cibles.get("commentsExtended"), plafond)
    commentaires = []
    for element in lire_partie(archive, nom, plafond):
        if nom_local(element.tag) != "comment":
            continue
        paragraphes = [p for p in element.iter() if nom_local(p.tag) == "p"]
        texte = "\n".join(textes_paragraphe(p)[0] for p in paragraphes)
        dernier = attribut_local(paragraphes[-1], "paraId") if paragraphes else None
        commentaires.append({
            "id": attribut_local(element, "id") or "?",
            "auteur": attribut_local(element, "author"),
            "date": attribut_local(element, "date"),
            "texte": extrait(texte),
            "resolu": resolus.get(dernier) if resolus is not None else None,
        })
    return commentaires


def etats_resolus(archive: zipfile.ZipFile, nom: str | None, plafond: int) -> dict[str | None, bool] | None:
    """paraId -> résolu, d'après w15:commentEx/@done ; None sans commentsExtended."""
    if nom is None or nom not in archive.namelist():
        return None
    return {
        attribut_local(e, "paraId"): attribut_local(e, "done") in ("1", "true")
        for e in lire_partie(archive, nom, plafond).iter() if nom_local(e.tag) == "commentEx"
    }


def suivi_active(archive: zipfile.ZipFile, nom: str | None, plafond: int) -> bool | None:
    """Le suivi des modifications est-il activé (w:trackRevisions) ?"""
    if nom is None or nom not in archive.namelist():
        return None
    for element in lire_partie(archive, nom, plafond):
        if nom_local(element.tag) == "trackRevisions":
            return attribut_local(element, "val") not in ("0", "false", "off")
    return False


def lire_word(chemin: Path, reglages: Reglages) -> Releve:
    """Moteur stdlib pour .docx/.docm/.dotx/.dotm."""
    releve = Releve(format=chemin.suffix.lower().lstrip("."))
    with ouvrir_archive(chemin) as archive:
        principale = partie_principale(archive, reglages.max_octets, "word/document.xml")
        rels = relations(archive, principale, reglages.max_octets)
        cibles = {type_rel: cible for _i, type_rel, cible in rels}
        styles = lire_styles(lire_si_present(archive, cibles.get("styles"), reglages.max_octets))
        releve.commentaires = lire_commentaires_word(archive, cibles, reglages.max_octets)
        releve.suivi_active = suivi_active(archive, cibles.get("settings"), reglages.max_octets)
        lire_corps_word(lire_partie(archive, principale, reglages.max_octets), styles, releve, principale)
        lire_parties_annexes(archive, rels, reglages.max_octets, releve)
        releve.parties.insert(0, principale)
    return releve


def lire_si_present(archive: zipfile.ZipFile, nom: str | None, plafond: int) -> ET.Element | None:
    """Partie XML si elle existe dans l'archive, sinon None."""
    if nom is None or nom not in archive.namelist():
        return None
    return lire_partie(archive, nom, plafond)


def lire_parties_annexes(archive: zipfile.ZipFile, rels: list[tuple[str, str, str]], plafond: int, releve: Releve) -> None:
    """Révisions des en-têtes, pieds de page et notes."""
    for _ident, type_rel, cible in rels:
        if type_rel in ("header", "footer", "footnotes", "endnotes") and cible in archive.namelist():
            releve.revisions.extend(revisions_word(lire_partie(archive, cible, plafond), cible, {}))
            releve.parties.append(cible)


# --------------------------------------------------------------------------- #
# PresentationML

def lire_powerpoint(chemin: Path, reglages: Reglages) -> Releve:
    """Moteur stdlib pour .pptx : diapositives dans l'ordre de présentation."""
    releve = Releve(format=chemin.suffix.lower().lstrip("."))
    with ouvrir_archive(chemin) as archive:
        principale = partie_principale(archive, reglages.max_octets, "ppt/presentation.xml")
        for numero, partie in enumerate(parties_diapositives(archive, principale, reglages.max_octets), start=1):
            racine = lire_partie(archive, partie, reglages.max_octets)
            diapositive = decrire_diapositive(racine, numero)
            ajouter_tableaux(releve, tableaux_drawingml(racine), numero)
            diapositive["notes"] = notes_diapositive(archive, partie, reglages.max_octets)
            releve.diapositives.append(diapositive)
            releve.parties.append(partie)
        ajouter_commentaires_pptx(archive, principale, releve, reglages.max_octets)
    remplir_paragraphes_diapositives(releve)
    return releve


def parties_diapositives(archive: zipfile.ZipFile, principale: str, plafond: int) -> list[str]:
    """Parties des diapositives dans l'ordre de p:sldIdLst."""
    cibles = {ident: cible for ident, _t, cible in relations(archive, principale, plafond)}
    ordre = []
    for element in lire_partie(archive, principale, plafond).iter():
        if nom_local(element.tag) == "sldId":
            cible = cibles.get(attribut_relation(element) or "")
            if cible and cible in archive.namelist():
                ordre.append(cible)
    return ordre


def tableaux_drawingml(racine: Any) -> list[dict[str, Any]]:
    """Tableaux a:tbl d'une diapositive : dimensions et première ligne."""
    tableaux = []
    for tableau in (e for e in racine.iter() if nom_local(e.tag) == "tbl"):
        lignes = [e for e in tableau if nom_local(e.tag) == "tr"]
        colonnes = sum(1 for e in tableau.iter() if nom_local(e.tag) == "gridCol")
        premiere = [extrait(" ".join(textes_drawingml(c))) for c in lignes[0] if nom_local(c.tag) == "tc"] if lignes else []
        tableaux.append({"lignes": len(lignes), "colonnes": colonnes, "premiere_ligne": premiere})
    return tableaux


def decrire_diapositive(racine: Any, numero: int) -> dict[str, Any]:
    """Titre, textes (a:p) et état masqué d'une diapositive."""
    titre = ""
    for forme in racine.iter():
        if nom_local(forme.tag) == "sp" and type_espace_reserve(forme) in ("title", "ctrTitle"):
            titre = "\n".join(textes_drawingml(forme))
            break
    return {
        "n": numero,
        "masquee": racine.get("show") == "0",
        "titre": titre,
        "textes": [t for t in textes_drawingml(racine) if t.strip()],
        "commentaires": [],
    }


def type_espace_reserve(forme: Any) -> str | None:
    """Type de p:ph d'une forme (title, body...), None si ce n'est pas un espace réservé."""
    for element in forme.iter():
        if nom_local(element.tag) == "ph":
            return element.get("type", "body")
    return None


def textes_drawingml(racine: Any) -> list[str]:
    """Texte de chaque a:p (runs, champs, sauts de ligne)."""
    textes = []
    for paragraphe in racine.iter():
        if nom_local(paragraphe.tag) != "p" or not paragraphe.tag.startswith("{http://schemas.openxmlformats.org/drawingml"):
            continue
        morceaux = []
        for element in paragraphe.iter():
            local = nom_local(element.tag)
            if local == "t":
                morceaux.append(element.text or "")
            elif local == "br":
                morceaux.append("\n")
        textes.append("".join(morceaux))
    return textes


def ajouter_tableaux(releve: Releve, tableaux: list[dict[str, Any]], numero: int) -> None:
    """Numérote les tableaux d'une diapositive à la suite des précédents."""
    for tableau in tableaux:
        releve.tableaux.append({"index": len(releve.tableaux) + 1, "diapositive": numero, **tableau})


def notes_diapositive(archive: zipfile.ZipFile, partie: str, plafond: int) -> str:
    """Texte des notes de l'orateur (espace réservé « body » du notesSlide)."""
    for _ident, type_rel, cible in relations(archive, partie, plafond):
        if type_rel == "notesSlide" and cible in archive.namelist():
            racine = lire_partie(archive, cible, plafond)
            formes = [f for f in racine.iter() if nom_local(f.tag) == "sp" and type_espace_reserve(f) == "body"]
            return "\n".join(t for f in formes for t in textes_drawingml(f)).strip()
    return ""


def ajouter_commentaires_pptx(archive: zipfile.ZipFile, principale: str, releve: Releve, plafond: int) -> None:
    """Commentaires anciens (p:cm) et modernes (p188:cm) rattachés à chaque diapositive."""
    auteurs = auteurs_pptx(archive, principale, plafond)
    for diapositive, partie in zip(releve.diapositives, releve.parties):
        for _ident, type_rel, cible in relations(archive, partie, plafond):
            if type_rel == "comments" and cible in archive.namelist():
                for element in lire_partie(archive, cible, plafond).iter():
                    if nom_local(element.tag) == "cm":
                        commentaire = commentaire_pptx(element, auteurs, diapositive["n"])
                        diapositive["commentaires"].append(commentaire)
                        releve.commentaires.append(commentaire)


def auteurs_pptx(archive: zipfile.ZipFile, principale: str, plafond: int) -> dict[str, str]:
    """Identifiant -> nom des auteurs de commentaires (commentAuthors.xml, authors.xml)."""
    auteurs = {}
    for _ident, type_rel, cible in relations(archive, principale, plafond):
        if type_rel in ("commentAuthors", "authors") and cible in archive.namelist():
            for element in lire_partie(archive, cible, plafond).iter():
                if nom_local(element.tag) in ("cmAuthor", "author"):
                    auteurs[element.get("id", "")] = element.get("name", "")
    return auteurs


def commentaire_pptx(element: Any, auteurs: dict[str, str], numero: int) -> dict[str, Any]:
    """Un commentaire PowerPoint, réponses comprises dans son texte."""
    morceaux = [e.text or "" for e in element.iter() if nom_local(e.tag) in ("text", "t")]
    auteur = element.get("authorId", "")
    return {
        "id": element.get("idx") or element.get("id") or "?",
        "auteur": auteurs.get(auteur, auteur or None),
        "date": element.get("dt") or element.get("created"),
        "texte": extrait(" ".join(morceaux)),
        "diapositive": numero,
        "reponses": sum(1 for e in element.iter() if nom_local(e.tag) == "reply"),
    }


def remplir_paragraphes_diapositives(releve: Releve) -> None:
    """Aplatit les textes des diapositives en paragraphes numérotés."""
    for diapositive in releve.diapositives:
        for texte in diapositive["textes"]:
            niveau = 1 if texte.strip() == diapositive["titre"] and texte.strip() else None
            releve.ajouter_paragraphe(texte, diapositive=diapositive["n"], niveau=niveau)


# --------------------------------------------------------------------------- #
# OpenDocument

@dataclass
class EtatOdf:
    """Insertions en cours (text:change-start) et leur texte."""

    actives: set[str] = field(default_factory=set)
    inseres: dict[str, list[str]] = field(default_factory=dict)
    commentaires: list[dict[str, Any]] = field(default_factory=list)


def lire_odf(chemin: Path, reglages: Reglages) -> Releve:
    """Moteur stdlib pour .odt/.ott/.odp/.otp."""
    with ouvrir_archive(chemin) as archive:
        if "content.xml" not in archive.namelist():
            raise ErreurDocument(f"{chemin.name} : content.xml absent, ce n'est pas un document OpenDocument")
        racine = lire_partie(archive, "content.xml", reglages.max_octets)
    releve = analyser_contenu_odf(racine, chemin.suffix.lower().lstrip("."))
    releve.parties.append("content.xml")
    return releve


def lire_odf_plat(chemin: Path, reglages: Reglages) -> Releve:
    """OpenDocument plat (.fodt/.fodp) : un seul fichier XML, lu borné."""
    with chemin.open("rb") as flux:
        donnees = flux.read(reglages.max_octets + 1)
    if len(donnees) > reglages.max_octets:
        raise ErreurDocument(f"{chemin.name} dépasse {reglages.max_octets} octets (voir --max-octets)", 1)
    releve = analyser_contenu_odf(analyser_xml(donnees, chemin.name), chemin.suffix.lower().lstrip("."))
    releve.parties.append(chemin.name)
    return releve


def analyser_contenu_odf(racine: ET.Element, format_doc: str) -> Releve:
    """Texte ou présentation selon le contenu de office:body."""
    releve = Releve(format=format_doc)
    corps = racine.find(f"{{{NS_OFFICE}}}body")
    if corps is None:
        raise ErreurDocument("office:body absent du contenu OpenDocument")
    texte = corps.find(f"{{{NS_OFFICE}}}text")
    presentation = corps.find(f"{{{NS_OFFICE}}}presentation")
    if texte is not None:
        lire_texte_odf(texte, releve)
    elif presentation is not None:
        lire_presentation_odf(presentation, masques_odp(racine), releve)
    else:
        raise ErreurDocument("ni office:text ni office:presentation : tableur ou dessin, hors domaine")
    return releve


def lire_texte_odf(texte: ET.Element, releve: Releve) -> None:
    """Paragraphes, titres, tableaux, annotations puis régions de révision."""
    etat = EtatOdf()
    for bloc in texte:
        traiter_bloc_odf(bloc, releve, etat, None)
    releve.commentaires.extend(etat.commentaires)
    suivies = texte.find(f"{{{NS_TEXT}}}tracked-changes")
    if suivies is not None:
        releve.revisions.extend(revisions_odf(suivies, etat))


def traiter_bloc_odf(bloc: ET.Element, releve: Releve, etat: EtatOdf, tableau: int | None) -> None:
    """Un élément de niveau bloc : paragraphe, titre, liste, section, tableau."""
    local = nom_local(bloc.tag)
    details: dict[str, Any] = {} if tableau is None else {"tableau": tableau}
    if bloc.tag == f"{{{NS_TEXT}}}h":
        niveau = int(bloc.get(f"{{{NS_TEXT}}}outline-level") or "1")
        releve.ajouter_paragraphe(texte_odf(bloc, etat), niveau=niveau, **details)
    elif bloc.tag == f"{{{NS_TEXT}}}p":
        releve.ajouter_paragraphe(texte_odf(bloc, etat), niveau=None, **details)
    elif bloc.tag == f"{{{NS_TABLE}}}table":
        traiter_tableau_odf(bloc, releve, etat)
    elif local in ("list", "list-item", "list-header", "section", "table-header-rows", "table-rows",
                   "table-row-group", "index-body", "table-of-content", "alphabetical-index",
                   "user-index", "illustration-index", "bibliography"):
        for enfant in bloc:
            traiter_bloc_odf(enfant, releve, etat, tableau)


def traiter_tableau_odf(tableau: ET.Element, releve: Releve, etat: EtatOdf) -> None:
    """Dimensions, première ligne, puis paragraphes des cellules."""
    index = len(releve.tableaux) + 1
    fiche: dict[str, Any] = {"index": index}
    releve.tableaux.append(fiche)
    lignes = list(lignes_tableau_odf(tableau))
    premiere: list[str] = []
    colonnes = 0
    for numero, ligne in enumerate(lignes):
        cellules = [c for c in ligne if nom_local(c.tag) in ("table-cell", "covered-table-cell")]
        colonnes = max(colonnes, sum(int(c.get(f"{{{NS_TABLE}}}number-columns-repeated") or "1") for c in cellules))
        for cellule in cellules:
            debut = len(releve.paragraphes)
            for bloc in cellule:
                traiter_bloc_odf(bloc, releve, etat, index)
            if numero == 0:
                premiere.append(extrait(" ".join(p["texte"] for p in releve.paragraphes[debut:])))
    fiche.update(lignes=len(lignes), colonnes=colonnes, premiere_ligne=premiere)


def lignes_tableau_odf(conteneur: ET.Element) -> Iterator[ET.Element]:
    """Lignes d'un tableau ODF, sans descendre dans les tableaux imbriqués."""
    for enfant in conteneur:
        local = nom_local(enfant.tag)
        if local == "table-row":
            yield enfant
        elif local in ("table-header-rows", "table-rows", "table-row-group"):
            yield from lignes_tableau_odf(enfant)


def texte_odf(element: ET.Element, etat: EtatOdf) -> str:
    """Texte d'un paragraphe ODF ; relève annotations et texte inséré au passage."""
    morceaux: list[str] = []
    ajouter_odf(element.text or "", morceaux, etat)
    for enfant in element:
        traiter_inline_odf(enfant, morceaux, etat)
        ajouter_odf(enfant.tail or "", morceaux, etat)
    return "".join(morceaux)


def ajouter_odf(texte: str, morceaux: list[str], etat: EtatOdf) -> None:
    """Ajoute un morceau au paragraphe et aux insertions ouvertes."""
    if not texte:
        return
    morceaux.append(texte)
    for ident in etat.actives:
        etat.inseres.setdefault(ident, []).append(texte)


def traiter_inline_odf(enfant: ET.Element, morceaux: list[str], etat: EtatOdf) -> None:
    """Élément de niveau caractère : espaces, tabulations, marques, annotations."""
    local = nom_local(enfant.tag)
    if local == "s":
        ajouter_odf(" " * int(enfant.get(f"{{{NS_TEXT}}}c") or "1"), morceaux, etat)
    elif local == "tab":
        ajouter_odf("\t", morceaux, etat)
    elif local == "line-break":
        ajouter_odf("\n", morceaux, etat)
    elif local == "change-start":
        etat.actives.add(enfant.get(f"{{{NS_TEXT}}}change-id") or "")
    elif local == "change-end":
        etat.actives.discard(enfant.get(f"{{{NS_TEXT}}}change-id") or "")
    elif local == "annotation":
        etat.commentaires.append(annotation_odf(enfant))
    elif local not in ("note", "annotation-end", "tracked-changes"):
        ajouter_odf(texte_odf(enfant, etat), morceaux, EtatOdf())


def annotation_odf(element: ET.Element) -> dict[str, Any]:
    """Un commentaire ODF (office:annotation ou officeooo:annotation)."""
    paragraphes = [texte_odf(p, EtatOdf()) for p in element if nom_local(p.tag) == "p"]
    resolu = attribut_local(element, "resolved")
    return {
        "id": attribut_local(element, "name") or "?",
        "auteur": element.findtext(f"{{{NS_DC}}}creator"),
        "date": element.findtext(f"{{{NS_DC}}}date"),
        "texte": extrait("\n".join(paragraphes)),
        "resolu": None if resolu is None else resolu == "true",
    }


def revisions_odf(suivies: ET.Element, etat: EtatOdf) -> list[dict[str, Any]]:
    """Régions text:changed-region : type, auteur, date, texte."""
    revisions = []
    types = {"insertion": "insertion", "deletion": "suppression", "format-change": "mise_en_forme"}
    for region in suivies:
        ident = region.get(f"{{{NS_TEXT}}}id") or attribut_local(region, "id") or ""
        for changement in region:
            local = nom_local(changement.tag)
            if local not in types:
                continue
            if local == "deletion":
                texte = " ".join(texte_odf(p, EtatOdf()) for p in changement if nom_local(p.tag) in ("p", "h"))
            else:
                texte = "".join(etat.inseres.get(ident, []))
            revisions.append({
                "type": types[local], "balise": f"text:{local}",
                "auteur": changement.findtext(f".//{{{NS_DC}}}creator"),
                "date": changement.findtext(f".//{{{NS_DC}}}date"),
                "texte": extrait(texte), "partie": "content.xml", "paragraphe": None,
            })
    return revisions


def masques_odp(racine: ET.Element) -> set[str]:
    """Styles de page dont presentation:visibility vaut hidden."""
    masques = set()
    for style in racine.iter(f"{{{NS_STYLE}}}style"):
        for propriete in style:
            if propriete.get(f"{{{NS_PRESENTATION}}}visibility") == "hidden":
                masques.add(style.get(f"{{{NS_STYLE}}}name") or "")
    return masques


def lire_presentation_odf(presentation: ET.Element, masques: set[str], releve: Releve) -> None:
    """Pages, titres, notes et annotations d'une présentation ODF."""
    for numero, page in enumerate(presentation.iter(f"{{{NS_DRAW}}}page"), start=1):
        notes = page.find(f"{{{NS_PRESENTATION}}}notes")
        annotations = [e for e in page.iter() if nom_local(e.tag) == "annotation"]
        exclus = {id(e) for bloc in ([notes] if notes is not None else []) + annotations for e in bloc.iter()}
        diapositive = {
            "n": numero,
            "masquee": (page.get(f"{{{NS_DRAW}}}style-name") or "") in masques,
            "titre": titre_page_odp(page),
            "textes": [texte_odf(p, EtatOdf()) for p in page.iter(f"{{{NS_TEXT}}}p") if id(p) not in exclus],
            "notes": "\n".join(texte_odf(p, EtatOdf()) for p in notes.iter(f"{{{NS_TEXT}}}p")) if notes is not None else "",
            "commentaires": [],
        }
        for annotation in annotations:
            commentaire = {**annotation_odf(annotation), "diapositive": numero}
            diapositive["commentaires"].append(commentaire)
            releve.commentaires.append(commentaire)
        diapositive["textes"] = [t for t in diapositive["textes"] if t.strip()]
        releve.diapositives.append(diapositive)
    remplir_paragraphes_diapositives(releve)


def titre_page_odp(page: ET.Element) -> str:
    """Texte du cadre presentation:class="title"."""
    for cadre in page.iter(f"{{{NS_DRAW}}}frame"):
        if cadre.get(f"{{{NS_PRESENTATION}}}class") == "title":
            return " ".join(texte_odf(p, EtatOdf()) for p in cadre.iter(f"{{{NS_TEXT}}}p")).strip()
    return ""


# --------------------------------------------------------------------------- #
# Moteurs optionnels

def lire_word_docx(chemin: Path, reglages: Reglages) -> Releve:
    """Moteur python-docx : styles et commentaires par la bibliothèque ;
    texte et révisions par le parcours stdlib de son arbre lxml (paragraph.text
    perd les insertions suivies)."""
    releve = lire_word(chemin, reglages)
    document = docx.Document(str(chemin))
    styles_bibliotheque = [ParagrapheDocx(p, document).style.name for p in paragraphes_corps(document.element)]
    for paragraphe, nom in zip(releve.paragraphes, styles_bibliotheque):
        if nom is not None and not paragraphe.get("zone_de_texte"):
            paragraphe["style"] = nom
    if hasattr(document, "comments"):
        ids = {str(c.comment_id): c for c in document.comments}
        for commentaire in releve.commentaires:
            source = ids.get(commentaire["id"])
            if source is not None:
                commentaire["auteur"] = source.author
                commentaire["texte"] = extrait(source.text)
    releve.moteur = "python-docx"
    return releve


def paragraphes_corps(racine: Any) -> list[Any]:
    """Paragraphes du corps dans l'ordre où lire_corps_word les numérote."""
    corps = next((e for e in racine if nom_local(e.tag) == "body"), racine)
    ordre: list[Any] = []
    for bloc in blocs_word(corps):
        if nom_local(bloc.tag) == "p":
            ordre.append(bloc)
        else:
            ordre.extend(paragraphes_tableau(bloc))
    return ordre


def paragraphes_tableau(tableau: Any) -> list[Any]:
    """Paragraphes d'un tableau, tableaux imbriqués compris, dans l'ordre."""
    ordre: list[Any] = []
    for ligne in (e for e in tableau if nom_local(e.tag) == "tr"):
        for cellule in (c for c in ligne if nom_local(c.tag) == "tc"):
            for bloc in blocs_word(cellule):
                ordre.extend([bloc] if nom_local(bloc.tag) == "p" else paragraphes_tableau(bloc))
    return ordre


def lire_powerpoint_pptx(chemin: Path, reglages: Reglages) -> Releve:
    """Moteur python-pptx pour les formes et les notes ; commentaires par la stdlib."""
    presentation = pptx.Presentation(str(chemin))
    releve = Releve(format=chemin.suffix.lower().lstrip("."), moteur="python-pptx")
    for numero, diapositive in enumerate(presentation.slides, start=1):
        titre = diapositive.shapes.title.text_frame.text if diapositive.shapes.title is not None else ""
        notes = ""
        if diapositive.has_notes_slide and diapositive.notes_slide.notes_text_frame is not None:
            notes = diapositive.notes_slide.notes_text_frame.text
        releve.diapositives.append({
            "n": numero, "masquee": diapositive.element.get("show") == "0",
            "titre": titre.replace("\v", "\n"),
            "textes": [t for t in textes_formes(diapositive.shapes) if t.strip()],
            "notes": notes.replace("\v", "\n").strip(), "commentaires": [],
        })
        releve.parties.append(diapositive.part.partname.lstrip("/"))
        ajouter_tableaux(releve, tableaux_drawingml(diapositive.element), numero)
    with ouvrir_archive(chemin) as archive:
        principale = partie_principale(archive, reglages.max_octets, "ppt/presentation.xml")
        ajouter_commentaires_pptx(archive, principale, releve, reglages.max_octets)
    remplir_paragraphes_diapositives(releve)
    return releve


def textes_formes(formes: Any) -> list[str]:
    """Paragraphes des formes (groupes et tableaux compris) via python-pptx."""
    textes: list[str] = []
    for forme in formes:
        if hasattr(forme, "shapes"):
            textes.extend(textes_formes(forme.shapes))
        if getattr(forme, "has_text_frame", False):
            textes.extend(p.text.replace("\v", "\n") for p in forme.text_frame.paragraphs)
        if getattr(forme, "has_table", False):
            for ligne in forme.table.rows:
                for cellule in ligne.cells:
                    textes.extend(p.text.replace("\v", "\n") for p in cellule.text_frame.paragraphs)
    return textes


# --------------------------------------------------------------------------- #
# Aiguillage

def lire_document(chemin: Path, reglages: Reglages) -> Releve:
    """Choisit le lecteur selon l'extension et le moteur demandé."""
    suffixe = chemin.suffix.lower()
    if suffixe in EXTENSIONS_BINAIRES:
        raise ErreurDocument(f"format binaire {suffixe} non pris en charge (convertir en {suffixe}x)")
    if suffixe in EXTENSIONS_WORD:
        return lire_avec_repli(chemin, reglages, docx, lire_word_docx, lire_word, "python-docx")
    if suffixe in EXTENSIONS_POWERPOINT:
        return lire_avec_repli(chemin, reglages, pptx, lire_powerpoint_pptx, lire_powerpoint, "python-pptx")
    if suffixe in EXTENSIONS_ODF:
        return lire_odf(chemin, reglages)
    if suffixe in EXTENSIONS_ODF_PLAT:
        return lire_odf_plat(chemin, reglages)
    raise ErreurDocument(f"extension « {suffixe or '(aucune)'} » non prise en charge")


def lire_avec_repli(chemin: Path, reglages: Reglages, module: Any, avec: Any, sans: Any, nom: str) -> Releve:
    """Bibliothèque si présente et demandée ; repli stdlib annoncé si elle échoue."""
    if reglages.moteur not in ("auto", "bibliotheque") or module is None:
        if reglages.moteur == "bibliotheque" and module is None:
            raise ErreurDocument(f"--moteur bibliotheque demandé mais {nom} n'est pas installé")
        return sans(chemin, reglages)
    try:
        return avec(chemin, reglages)
    except ErreurDocument:
        raise
    except ERREURS_LECTURE + (TypeError, AttributeError, IndexError) as exc:
        print(f"{nom} a échoué sur {chemin.name} ({type(exc).__name__}: {exc}) ; repli stdlib", file=sys.stderr)
        releve = sans(chemin, reglages)
        releve.notes.append(f"{nom} : {type(exc).__name__}, lu par la stdlib")
        return releve


def lire_xml_en_ligne(texte: str, reglages: Reglages) -> Releve:
    """Analyse une partie XML passée en argument (document.xml, content.xml, slide)."""
    donnees = texte.encode("utf-8")
    if len(donnees) > reglages.max_octets:
        raise ErreurDocument(f"--xml dépasse {reglages.max_octets} octets")
    racine = analyser_xml(donnees, "--xml")
    local = nom_local(racine.tag)
    if local == "document" and "wordprocessingml" in racine.tag:
        releve = Releve(format="xml-word")
        lire_corps_word(racine, Styles(), releve, "(--xml)")
        return releve
    if local in ("document-content", "document"):
        return analyser_contenu_odf(racine, "xml-odf")
    if local == "sld":
        releve = Releve(format="xml-diapositive")
        releve.diapositives.append(decrire_diapositive(racine, 1) | {"notes": ""})
        remplir_paragraphes_diapositives(releve)
        return releve
    raise ErreurDocument(f"racine XML « {local} » non reconnue (attendu w:document, office:document-content ou p:sld)")


def lister_documents(dossier: Path, plafond: int) -> tuple[list[Path], list[str]]:
    """Documents d'un dossier (récursif, trié) ; fichiers verrous écartés."""
    acceptees = EXTENSIONS_WORD + EXTENSIONS_POWERPOINT + EXTENSIONS_ODF + EXTENSIONS_ODF_PLAT
    retenus, ecartes = [], []
    for chemin in sorted(dossier.rglob("*")):
        if not chemin.is_file() or chemin.suffix.lower() not in acceptees:
            continue
        if chemin.name.startswith(("~$", ".~lock")):
            ecartes.append(f"{chemin.relative_to(dossier)} (fichier verrou)")
        elif len(retenus) < plafond:
            retenus.append(chemin)
        else:
            ecartes.append(f"{chemin.relative_to(dossier)} (au-delà de --max-fichiers)")
    return retenus, ecartes


def examiner_chemin(cible: Path, reglages: Reglages, max_fichiers: int) -> dict[str, Any]:
    """Examine un document ou tous ceux d'un dossier."""
    if not cible.exists():
        raise ErreurDocument(f"chemin introuvable : {cible}")
    if cible.is_dir():
        fichiers, ecartes = lister_documents(cible, max_fichiers)
        base = cible
    else:
        fichiers, ecartes, base = [cible], [], cible.parent
    documents, illisibles = [], []
    for fichier in fichiers:
        nom = str(fichier.relative_to(base))
        try:
            documents.append(presenter_releve(nom, lire_document(fichier, reglages), reglages))
        except ERREURS_LECTURE as exc:
            if not cible.is_dir():
                raise ErreurDocument(f"{nom} : {exc}", getattr(exc, "code", 2)) from exc
            illisibles.append({"fichier": nom, "erreur": f"{type(exc).__name__}: {exc}"})
    return {"documents": documents, "illisibles": illisibles, "ecartes": ecartes}


def presenter_releve(nom: str, releve: Releve, reglages: Reglages) -> dict[str, Any]:
    """Dictionnaire publié pour un document."""
    textes = [p["texte"] for p in releve.paragraphes]
    return {
        "fichier": nom,
        "format": releve.format,
        "moteur": releve.moteur,
        "paragraphes_total": len(releve.paragraphes),
        "mots": sum(len(t.split()) for t in textes),
        "titres": releve.titres,
        "tableaux": releve.tableaux,
        "revisions": bilan_revisions(releve.revisions),
        "commentaires": {"total": len(releve.commentaires), "detail": releve.commentaires},
        "suivi_des_modifications_active": releve.suivi_active,
        "diapositives": releve.diapositives,
        "paragraphes": releve.paragraphes[: reglages.max_paragraphes],
        "paragraphes_tronques": len(releve.paragraphes) > reglages.max_paragraphes,
        "parties_examinees": releve.parties,
        "notes": releve.notes,
    }


def bilan_revisions(revisions: list[dict[str, Any]]) -> dict[str, Any]:
    """Comptes par type et par auteur, puis le détail."""
    return {
        "total": len(revisions),
        "par_type": dict(Counter(r["type"] for r in revisions).most_common()),
        "par_auteur": dict(Counter(r["auteur"] or "(inconnu)" for r in revisions).most_common()),
        "detail": revisions,
    }


def assembler(resultat: dict[str, Any]) -> dict[str, Any]:
    """Objet JSON final ; `denominateur` en tête (documents réellement examinés)."""
    documents = resultat["documents"]
    noms = [d["fichier"] for d in documents] + [i["fichier"] for i in resultat["illisibles"]]
    moteurs = sorted({d["moteur"] for d in documents}) or ["stdlib"]
    return {
        "denominateur": len(noms),
        "examines": noms[:LIMITE_EXAMINES],
        "examines_tronques": len(noms) > LIMITE_EXAMINES,
        "moteur": "+".join(moteurs),
        "revisions_non_acceptees": sum(d["revisions"]["total"] for d in documents),
        "commentaires": sum(d["commentaires"]["total"] for d in documents),
        "illisibles": resultat["illisibles"],
        "ecartes": resultat["ecartes"],
        "documents": documents,
        "contrat": extraire_contrat(__doc__ or ""),
    }


# --------------------------------------------------------------------------- #
# Présentation

def afficher_json(sortie: dict[str, Any]) -> None:
    """Un seul objet JSON sur stdout."""
    print(json.dumps(sortie, ensure_ascii=False, indent=2))


def afficher_humain(sortie: dict[str, Any], integral: bool) -> None:
    """Résumé lisible ; texte complet avec --integral."""
    print(f"{sortie['denominateur']} document(s) examiné(s), moteur {sortie['moteur']} ; "
          f"{sortie['revisions_non_acceptees']} révision(s) non acceptée(s), {sortie['commentaires']} commentaire(s)")
    for document in sortie["documents"]:
        afficher_document(document, integral)
    for illisible in sortie["illisibles"]:
        print(f"\nILLISIBLE {illisible['fichier']} : {illisible['erreur']}")


def afficher_document(document: dict[str, Any], integral: bool) -> None:
    """Bloc lisible d'un document."""
    print(f"\n{document['fichier']} ({document['format']}, moteur {document['moteur']}) : "
          f"{document['paragraphes_total']} paragraphe(s), {document['mots']} mot(s), "
          f"{len(document['tableaux'])} tableau(x)")
    if document["suivi_des_modifications_active"]:
        print("  suivi des modifications ACTIVÉ")
    for titre in document["titres"]:
        print(f"  {'  ' * max(titre['niveau'] - 1, 0)}# {titre['texte']}")
    for diapositive in document["diapositives"]:
        masque = " [masquée]" if diapositive["masquee"] else ""
        print(f"  diapositive {diapositive['n']}{masque} : {diapositive['titre'] or '(sans titre)'}"
              + (f" — notes : {extrait(diapositive['notes'])}" if diapositive["notes"] else ""))
    for revision in document["revisions"]["detail"]:
        print(f"  RÉVISION {revision['type']} par {revision['auteur'] or '?'} ({revision['date'] or 'sans date'})"
              f" §{revision['paragraphe'] or '-'} : {revision['texte']!r}")
    for commentaire in document["commentaires"]["detail"]:
        etat = {True: " [résolu]", False: " [ouvert]"}.get(commentaire.get("resolu"), "")
        print(f"  COMMENTAIRE {commentaire['auteur'] or '?'}{etat} : {commentaire['texte']}")
    if integral:
        for paragraphe in document["paragraphes"]:
            print(f"  {paragraphe['n']:>4} | {paragraphe['texte']}")


# --------------------------------------------------------------------------- #
# Interface

def construire_analyseur() -> argparse.ArgumentParser:
    """Analyseur d'arguments en français."""
    analyseur = argparse.ArgumentParser(
        description="Lit un document Word, PowerPoint ou OpenDocument (texte, titres, tableaux, "
                    "diapositives, notes) et signale les révisions non acceptées et les "
                    "commentaires. Code 1 s'il y en a.",
        epilog="Exemple : python3 lire_bureautique.py livrables/contrat.docx --json",
    )
    analyseur.add_argument("chemin", nargs="?", help="document, ou dossier dont examiner les documents")
    analyseur.add_argument("--xml", help="contenu XML d'une partie (word/document.xml, content.xml, slide)")
    analyseur.add_argument("--json", action="store_true", help="un seul objet JSON sur stdout")
    analyseur.add_argument("--racine", help="dossier de base des chemins relatifs (défaut : dossier courant)")
    analyseur.add_argument("--moteur", choices=("auto", "stdlib", "bibliotheque"), default="auto",
                           help="auto : python-docx / python-pptx s'ils sont installés, stdlib sinon")
    analyseur.add_argument("--integral", action="store_true", help="afficher tout le texte (sortie lisible)")
    analyseur.add_argument("--max-paragraphes", type=int, default=5000, help="paragraphes listés en JSON par document")
    analyseur.add_argument("--max-octets", type=int, default=64 * 1024 * 1024,
                           help="plafond d'octets décompressés par partie XML (garde contre les bombes zip)")
    analyseur.add_argument("--max-fichiers", type=int, default=500, help="plafond de documents examinés dans un dossier")
    return analyseur


def resoudre(chemin: str, racine: str | None) -> Path:
    """Chemin relatif résolu contre --racine (ou le dossier courant)."""
    brut = Path(chemin)
    if brut.is_absolute() or racine is None:
        return brut
    return Path(racine) / brut


def annoncer_replis(moteur: str, sortie: dict[str, Any]) -> None:
    """Une ligne sur stderr quand une bibliothèque utile manque."""
    formats = {d["format"] for d in sortie["documents"]}
    manquantes = []
    if docx is None and formats & {e.lstrip(".") for e in EXTENSIONS_WORD}:
        manquantes.append("python-docx")
    if pptx is None and formats & {e.lstrip(".") for e in EXTENSIONS_POWERPOINT}:
        manquantes.append("python-pptx")
    if manquantes and moteur == "auto":
        print(f"{' et '.join(manquantes)} absent(s) : lecture en mode dégradé stdlib (zipfile + xml.etree)",
              file=sys.stderr)


def main() -> int:
    """Point d'entrée : 0 rien à signaler, 1 révision ou commentaire, 2 entrée invalide, 3 rien à examiner."""
    args = construire_analyseur().parse_args()
    if (args.chemin is None) == (args.xml is None):
        print("donner soit un chemin, soit --xml (et pas les deux)", file=sys.stderr)
        return 2
    reglages = Reglages(args.moteur, max(args.max_octets, 1024), max(args.max_paragraphes, 0))
    try:
        if args.xml is not None:
            resultat = {"documents": [presenter_releve("(--xml)", lire_xml_en_ligne(args.xml, reglages), reglages)],
                        "illisibles": [], "ecartes": []}
        else:
            resultat = examiner_chemin(resoudre(args.chemin, args.racine), reglages, max(args.max_fichiers, 1))
    except ErreurDocument as exc:
        print(f"lire_bureautique : {exc}", file=sys.stderr)
        if args.json:
            afficher_json({"denominateur": 0, "examines": [], "moteur": "stdlib", "erreur": str(exc), "code": exc.code})
        return exc.code
    sortie = assembler(resultat)
    annoncer_replis(args.moteur, sortie)
    if args.json:
        afficher_json(sortie)
    else:
        afficher_humain(sortie, args.integral)
    if sortie["denominateur"] == 0:
        print("dénominateur nul : aucun document trouvé, rien à examiner", file=sys.stderr)
        return 3
    defaut = sortie["revisions_non_acceptees"] or sortie["commentaires"] or sortie["illisibles"]
    return 1 if defaut else 0


if __name__ == "__main__":
    raise SystemExit(main())
