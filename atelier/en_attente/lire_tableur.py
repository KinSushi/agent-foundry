"""Un classeur n'est pas du texte : un .xlsx est une archive zip de parties XML, et
l'interpréteur propre de cette boîte n'a aucun lecteur (mesuré : `python3.14 -c
'import openpyxl'` rend ModuleNotFoundError). Le modèle default.ots livré avec
LibreOffice 24.2 déclare une ligne vide répétée 1016575 fois (mesuré : `unzip -p
default.ots content.xml | grep -o 'number-rows-repeated="[0-9]*"'`) : un lecteur
qui déplie les répétitions fabrique un million de lignes vides.

QUESTION
    Que contient ce classeur : feuilles, dimensions, en-têtes, types par colonne,
    cellules en erreur, formules ?
MESURE
    Lecture en flux des parties XML (zipfile + xml.etree.iterparse) ou du texte
    délimité (module csv, délimiteur deviné par csv.Sniffer). Pour chaque feuille :
    les cellules non vides, leur type (texte, nombre, date d'après le format
    numérique du style, booléen, erreur, texte qui ressemble à un nombre), les
    valeurs d'erreur mises en cache (t="e" en Office Open XML,
    calcext:value-type="error" en OpenDocument), le texte des formules et les
    formules sans valeur calculée. Moteur openpyxl (lecture seule, dimensions
    réinitialisées) pour .xlsx/.xlsm s'il est installé ; python-calamine pour
    .xls/.xlsb, que la stdlib ne sait pas lire.
HYPOTHÈSES
    Les valeurs en cache reflètent le dernier recalcul : l'outil ne recalcule rien.
    La première ligne non vide d'une feuille est sa ligne d'en-têtes. Un .csv/.tsv
    est en utf-8 (avec ou sans BOM) ou, à défaut, en cp1252.
LIMITES
    N'évalue aucune formule : une formule jamais recalculée (fichier écrit par une
    bibliothèque) n'a pas de valeur, son erreur éventuelle est invisible ; elle est
    comptée dans « sans_valeur ». Sans python-calamine, .xls/.xlsb sont refusés ;
    avec python-calamine, les cellules en erreur sont lues comme vides (mesuré sur
    python-calamine 0.8.2) et les formules sont inaccessibles : le rapport le dit.
    Le moteur stdlib ne retranscrit pas les formules partagées (marquées
    « partagée »). Chiffrement, tableaux croisés, graphiques, macros, mises en forme
    conditionnelles : ignorés. Lecture plafonnée (--max-cellules par feuille,
    --max-octets par partie décompressée) : au-delà, la feuille est dite tronquée.
CONTRE-EXEMPLES
    Sur cl-test.ods (livré avec LibreOffice 24.2), la ligne 1 porte une note
    (« Max », « We paste a =MAX(results)... ») et l'en-tête du tableau est en ligne
    4 : l'outil rend la note comme en-têtes et les types par colonne mêlent la note
    et les données.
INVOCATION
    {outil} {dossier}/echantillon.csv --json
DOMAINE
    Classeurs .xlsx/.xlsm/.xltx/.xltm, .ods/.ots, .csv/.tsv (et .xls/.xlsb avec
    python-calamine), lus sans rien exécuter ni recalculer, pour décider quoi
    extraire ou quoi corriger avant un traitement automatique.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import re
import sys
import zipfile
import zlib
import xml.etree.ElementTree as ET
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterator, NamedTuple

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

try:
    import openpyxl  # type: ignore[import-not-found]
except ImportError:
    openpyxl = None

try:
    import python_calamine  # type: ignore[import-not-found]
except ImportError:
    python_calamine = None

RACINE = Path(__file__).resolve().parent

INTITULE_QUESTION = "QUESTION"
INTITULE_MESURE = "MESURE"
INTITULE_HYPOTHESES = "HYPOTHÈSES"
INTITULE_LIMITES = "LIMITES"
INTITULE_CONTRE_EXEMPLES = "CONTRE-EXEMPLES"
INTITULE_INVOCATION = "INVOCATION"
INTITULE_DOMAINE = "DOMAINE"
INTITULES = (
    INTITULE_QUESTION, INTITULE_MESURE, INTITULE_HYPOTHESES, INTITULE_LIMITES,
    INTITULE_CONTRE_EXEMPLES, INTITULE_INVOCATION, INTITULE_DOMAINE,
)

EXTENSIONS_OOXML = (".xlsx", ".xlsm", ".xltx", ".xltm")
EXTENSIONS_ODF = (".ods", ".ots")
EXTENSIONS_TEXTE = (".csv", ".tsv")
EXTENSIONS_BINAIRES = (".xls", ".xlsb")
CODES_ERREUR = frozenset({
    "#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A",
    "#GETTING_DATA", "#SPILL!", "#CALC!", "#FIELD!", "#BLOCKED!", "#CONNECT!",
    "#BUSY!", "#UNKNOWN!", "#EXTERNAL!",
})
MOTIF_ERREUR_LIBREOFFICE = re.compile(r"Err:\d{3}")
MOTIF_NOMBRE = re.compile(r"[-+]?(\d+([.,]\d*)?|[.,]\d+)([eE][-+]?\d+)?")
MOTIF_DATE_TEXTE = re.compile(r"\d{4}-\d{2}-\d{2}([T ]\d{2}:\d{2}(:\d{2}(\.\d+)?)?)?|\d{1,2}/\d{1,2}/\d{4}")
MOTIF_REFERENCE = re.compile(r"([A-Z]+)(\d+)")
BOOLEENS_TEXTE = frozenset({"true", "false", "vrai", "faux"})
FORMATS_DATE_INTEGRES = frozenset(
    list(range(14, 23)) + list(range(27, 37)) + [45, 46, 47] + list(range(50, 59))
)
SIGNATURE_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ORIGINE_1900 = dt.datetime(1899, 12, 30)
ORIGINE_1904 = dt.datetime(1904, 1, 1)
LIMITE_EXEMPLES = 20
LIMITE_EXAMINES = 200
LIGNE_MAX = 1 << 20
CANDIDATS_DELIMITEUR = ",;\t|"
NS_OFFICE = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
NS_TABLE = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
NS_TEXT = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
NS_STYLE = "urn:oasis:names:tc:opendocument:xmlns:style:1.0"
NS_CALCEXT = "urn:org:documentfoundation:names:experimental:calc:xmlns:calcext:1.0"


class ErreurClasseur(Exception):
    """Entrée illisible ou refusée ; `code` est le code de sortie proposé."""

    def __init__(self, message: str, code: int = 2) -> None:
        super().__init__(message)
        self.code = code


ERREURS_LECTURE = (
    OSError, zipfile.BadZipFile, zipfile.LargeZipFile, ET.ParseError, KeyError,
    ValueError, csv.Error, EOFError, zlib.error, ErreurClasseur,
)


class Cellule(NamedTuple):
    """Une cellule non vide, normalisée quel que soit le format d'origine."""

    ligne: int
    colonne: int
    genre: str
    valeur: Any
    formule: str | None


@dataclass
class Bilan:
    """Accumule les cellules d'une feuille et en tire le rapport."""

    nom: str
    etat: str
    max_cellules: int
    apercu_max: int
    formules_connues: bool = True
    erreurs_connues: bool = True
    ligne_entete: int | None = None
    entetes: dict[int, str] = field(default_factory=dict)
    lignes: tuple[int, int] = (0, 0)
    colonnes: tuple[int, int] = (0, 0)
    non_vides: int = 0
    types: dict[int, Counter] = field(default_factory=dict)
    erreurs: list[dict[str, Any]] = field(default_factory=list)
    nb_erreurs: int = 0
    nb_formules: int = 0
    sans_valeur: int = 0
    exemples_formules: list[dict[str, str]] = field(default_factory=list)
    apercu: dict[int, dict[int, Any]] = field(default_factory=dict)
    tronque: bool = False
    notes: list[str] = field(default_factory=list)

    def ajouter(self, cellule: Cellule) -> bool:
        """Intègre une cellule ; rend False quand le plafond est atteint."""
        if self.non_vides >= self.max_cellules:
            self.tronque = True
            return False
        self.non_vides += 1
        self._etendre(cellule.ligne, cellule.colonne)
        if self.ligne_entete is None:
            self.ligne_entete = cellule.ligne
        if cellule.ligne == self.ligne_entete:
            self.entetes[cellule.colonne] = texte_court(valeur_json(cellule.valeur))
        else:
            self.types.setdefault(cellule.colonne, Counter())[cellule.genre] += 1
            self._retenir_apercu(cellule)
        self._retenir_formule(cellule)
        self._retenir_erreur(cellule)
        return True

    def _etendre(self, ligne: int, colonne: int) -> None:
        if self.non_vides == 1:
            self.lignes, self.colonnes = (ligne, ligne), (colonne, colonne)
            return
        self.lignes = (min(self.lignes[0], ligne), max(self.lignes[1], ligne))
        self.colonnes = (min(self.colonnes[0], colonne), max(self.colonnes[1], colonne))

    def _retenir_apercu(self, cellule: Cellule) -> None:
        if cellule.ligne in self.apercu or len(self.apercu) < self.apercu_max:
            self.apercu.setdefault(cellule.ligne, {})[cellule.colonne] = valeur_json(cellule.valeur)

    def _retenir_formule(self, cellule: Cellule) -> None:
        if cellule.formule is None:
            return
        self.nb_formules += 1
        if cellule.genre == "sans_valeur":
            self.sans_valeur += 1
        if len(self.exemples_formules) < LIMITE_EXEMPLES:
            self.exemples_formules.append(
                {"cellule": reference(cellule.ligne, cellule.colonne), "formule": cellule.formule}
            )

    def _retenir_erreur(self, cellule: Cellule) -> None:
        if cellule.genre != "erreur":
            return
        self.nb_erreurs += 1
        if len(self.erreurs) < LIMITE_EXEMPLES:
            self.erreurs.append({
                "cellule": reference(cellule.ligne, cellule.colonne),
                "valeur": str(cellule.valeur),
                "formule": cellule.formule,
            })

    def rapport(self) -> dict[str, Any]:
        """Rend le dictionnaire publié pour cette feuille."""
        vide = self.non_vides == 0
        return {
            "nom": self.nom,
            "etat": self.etat,
            "dimension": None if vide else plage(self.lignes, self.colonnes),
            "lignes": 0 if vide else self.lignes[1] - self.lignes[0] + 1,
            "colonnes": 0 if vide else self.colonnes[1] - self.colonnes[0] + 1,
            "cellules_non_vides": self.non_vides,
            "ligne_entete": self.ligne_entete,
            "entetes": [self.entetes.get(c, "") for c in self._plage_colonnes()],
            "colonnes_detail": self._detail_colonnes(),
            "nb_erreurs": self.nb_erreurs if self.erreurs_connues else None,
            "erreurs": self.erreurs if self.erreurs_connues else None,
            "nb_formules": self.nb_formules if self.formules_connues else None,
            "formules_sans_valeur": self.sans_valeur if self.formules_connues else None,
            "exemples_formules": self.exemples_formules if self.formules_connues else None,
            "apercu": [self._ligne_apercu(n) for n in sorted(self.apercu)],
            "tronque": self.tronque,
            "notes": self.notes,
        }

    def _plage_colonnes(self) -> range:
        if self.non_vides == 0:
            return range(0)
        return range(self.colonnes[0], self.colonnes[1] + 1)

    def _ligne_apercu(self, numero: int) -> dict[str, Any]:
        cellules = self.apercu[numero]
        return {"ligne": numero, "valeurs": [cellules.get(c) for c in self._plage_colonnes()]}

    def _detail_colonnes(self) -> list[dict[str, Any]]:
        lignes_donnees = 0
        if self.ligne_entete is not None and self.non_vides:
            lignes_donnees = self.lignes[1] - self.ligne_entete
        detail = []
        for colonne in self._plage_colonnes():
            compte = self.types.get(colonne, Counter())
            remplies = sum(compte.values())
            detail.append({
                "colonne": lettres_colonne(colonne),
                "entete": self.entetes.get(colonne, ""),
                "types": dict(compte.most_common()),
                "vides": max(lignes_donnees - remplies, 0),
                "mixte": len([g for g in compte if g != "sans_valeur"]) > 1,
            })
        return detail


@dataclass
class Reglages:
    """Paramètres de lecture communs à tous les formats."""

    moteur: str
    max_cellules: int
    max_octets: int
    apercu: int
    delimiteur: str | None


# --------------------------------------------------------------------------- #
# Outils de base

def extraire_contrat(doc: str) -> dict[str, str]:
    """Découpe la docstring en sections selon les intitulés du contrat."""
    contrat: dict[str, str] = {}
    courant: str | None = None
    for ligne in doc.splitlines():
        tete = ligne.strip()
        if tete in INTITULES:
            courant = tete
            contrat[courant] = ""
        elif courant is not None:
            contrat[courant] = (contrat[courant] + " " + tete).strip()
    return contrat


def lettres_colonne(numero: int) -> str:
    """1 -> A, 27 -> AA."""
    lettres = ""
    while numero > 0:
        numero, reste = divmod(numero - 1, 26)
        lettres = chr(65 + reste) + lettres
    return lettres


def numero_colonne(lettres: str) -> int:
    """A -> 1, AA -> 27."""
    numero = 0
    for lettre in lettres:
        numero = numero * 26 + (ord(lettre) - 64)
    return numero


def reference(ligne: int, colonne: int) -> str:
    """(3, 4) -> D3."""
    return f"{lettres_colonne(colonne)}{ligne}"


def plage(lignes: tuple[int, int], colonnes: tuple[int, int]) -> str:
    """Plage A1:C9 couvrant les cellules non vides."""
    return f"{reference(lignes[0], colonnes[0])}:{reference(lignes[1], colonnes[1])}"


def texte_court(valeur: Any, limite: int = 80) -> str:
    """Représentation textuelle bornée d'une valeur de cellule."""
    if valeur is None:
        return ""
    texte = valeur if isinstance(valeur, str) else str(valeur)
    return texte if len(texte) <= limite else texte[: limite - 1] + "…"


def valeur_json(valeur: Any) -> Any:
    """Valeur sérialisable et bornée pour l'aperçu."""
    if valeur is None or isinstance(valeur, (bool, int)):
        return valeur
    if isinstance(valeur, float):
        return int(valeur) if valeur.is_integer() and abs(valeur) < 2**53 else valeur
    return texte_court(valeur)


def nom_local(etiquette: str) -> str:
    """Retire l'espace de noms {uri} d'une étiquette ElementTree."""
    return etiquette.rsplit("}", 1)[-1]


def attribut_relation(element: Any) -> str | None:
    """Attribut r:id (espace de noms des relations, transitionnel ou strict)."""
    for cle, valeur in element.attrib.items():
        if cle.endswith("}id") and "relationships" in cle:
            return valeur
    return None


def genre_texte(texte: str) -> str:
    """Distingue un texte ordinaire d'un nombre stocké comme texte."""
    propre = texte.strip()
    if propre and MOTIF_NOMBRE.fullmatch(propre):
        return "texte_numerique"
    return "texte"


def est_code_erreur(texte: str) -> bool:
    """Vrai pour #DIV/0!, #N/A, Err:502..."""
    propre = texte.strip()
    return propre in CODES_ERREUR or bool(MOTIF_ERREUR_LIBREOFFICE.fullmatch(propre))


def date_depuis_serie(serie: float, origine: dt.datetime) -> str | None:
    """Numéro de série Excel -> date ISO ; None si hors bornes."""
    if serie < 0 or serie > 2958465:
        return None
    if origine == ORIGINE_1900 and serie < 61:
        serie += 1
    moment = origine + dt.timedelta(days=serie)
    if serie < 1:
        return moment.time().isoformat(timespec="seconds")
    if moment.time() == dt.time(0, 0):
        return moment.date().isoformat()
    return moment.isoformat(timespec="seconds")


# --------------------------------------------------------------------------- #
# Lecture bornée des parties d'une archive

class LecteurBorne(io.RawIOBase):
    """Flux qui compte les octets décompressés, plafonne et refuse toute DTD."""

    def __init__(self, flux: Any, plafond: int, nom: str) -> None:
        super().__init__()
        self._flux = flux
        self._plafond = plafond
        self._nom = nom
        self._lus = 0
        self._queue = b""

    def readable(self) -> bool:
        return True

    def readinto(self, tampon: Any) -> int:
        donnees = self._flux.read(len(tampon))
        self._lus += len(donnees)
        if self._lus > self._plafond:
            raise ErreurClasseur(
                f"{self._nom} : plus de {self._plafond} octets décompressés, lecture "
                "arrêtée (garde contre les bombes zip ; voir --max-octets)", 1)
        fenetre = self._queue + donnees
        if b"<!DOCTYPE" in fenetre or b"<!ENTITY" in fenetre:
            raise ErreurClasseur(f"{self._nom} : déclaration DTD refusée (jamais légitime ici)", 1)
        self._queue = donnees[-16:]
        tampon[: len(donnees)] = donnees
        return len(donnees)


def ouvrir_partie(archive: zipfile.ZipFile, nom: str, plafond: int) -> LecteurBorne:
    """Ouvre une partie de l'archive après contrôle de sa taille déclarée."""
    info = archive.getinfo(nom)
    if info.file_size > plafond:
        raise ErreurClasseur(
            f"{nom} déclare {info.file_size} octets décompressés, au-delà du plafond "
            f"{plafond} (garde contre les bombes zip ; voir --max-octets)", 1)
    return LecteurBorne(archive.open(info), plafond, nom)


def verifier_tailles(archive: zipfile.ZipFile, plafond: int) -> None:
    """Refuse une archive dont une partie déclare plus que le plafond décompressé."""
    for info in archive.infolist():
        if info.file_size > plafond:
            raise ErreurClasseur(
                f"{info.filename} déclare {info.file_size} octets décompressés, au-delà du "
                f"plafond {plafond} (garde contre les bombes zip ; voir --max-octets)", 1)


def inventaire_ooxml(chemin: Path) -> dict[str, Any]:
    """Macros VBA et liaisons externes, visibles dans la liste des parties."""
    with ouvrir_archive(chemin) as archive:
        noms = archive.namelist()
    return {
        "macros_vba": any(n.lower().endswith("vbaproject.bin") for n in noms),
        "liaisons_externes": sum(1 for n in noms if n.startswith("xl/externalLinks/") and n.endswith(".xml")),
    }


def lire_xml(archive: zipfile.ZipFile, nom: str, plafond: int) -> ET.Element:
    """Lit entièrement une petite partie XML (classeur, styles, relations)."""
    with ouvrir_partie(archive, nom, plafond) as flux:
        return ET.parse(flux).getroot()


def ouvrir_archive(chemin: Path) -> zipfile.ZipFile:
    """Ouvre l'archive ou explique pourquoi ce n'en est pas une."""
    with chemin.open("rb") as brut:
        signature = brut.read(8)
    if signature == SIGNATURE_OLE:
        raise ErreurClasseur(
            f"{chemin.name} est un conteneur OLE (classeur chiffré ou .xls renommé), pas une archive zip")
    try:
        return zipfile.ZipFile(chemin)
    except zipfile.BadZipFile as exc:
        raise ErreurClasseur(f"{chemin.name} n'est pas une archive zip valide ({exc})") from exc


# --------------------------------------------------------------------------- #
# Office Open XML (stdlib)

def relations(archive: zipfile.ZipFile, partie: str, plafond: int) -> dict[str, tuple[str, str]]:
    """Identifiant -> (type, cible résolue) pour les relations d'une partie."""
    dossier, _, nom = partie.rpartition("/")
    chemin_rels = f"{dossier}/_rels/{nom}.rels" if dossier else f"_rels/{nom}.rels"
    if chemin_rels not in archive.namelist():
        return {}
    resultat = {}
    for rel in lire_xml(archive, chemin_rels, plafond):
        cible = rel.get("Target", "")
        if rel.get("TargetMode") == "External":
            continue
        resolue = cible.lstrip("/") if cible.startswith("/") else normaliser_chemin(dossier, cible)
        resultat[rel.get("Id", "")] = (rel.get("Type", "").rsplit("/", 1)[-1], resolue)
    return resultat


def normaliser_chemin(dossier: str, cible: str) -> str:
    """Résout ../ dans un chemin interne à l'archive."""
    morceaux = [m for m in dossier.split("/") if m]
    for morceau in cible.split("/"):
        if morceau == "..":
            if morceaux:
                morceaux.pop()
        elif morceau not in ("", "."):
            morceaux.append(morceau)
    return "/".join(morceaux)


def partie_classeur(archive: zipfile.ZipFile, plafond: int) -> str:
    """Chemin de xl/workbook.xml d'après _rels/.rels."""
    for type_rel, cible in relations(archive, "", plafond).values():
        if type_rel == "officeDocument":
            return cible
    return "xl/workbook.xml"


def chaines_partagees(archive: zipfile.ZipFile, nom: str | None, plafond: int) -> list[str]:
    """Table des chaînes partagées, lue en flux (texte phonétique exclu)."""
    if nom is None or nom not in archive.namelist():
        return []
    chaines: list[str] = []
    with ouvrir_partie(archive, nom, plafond) as flux:
        for _evenement, element in ET.iterparse(flux, events=("end",)):
            if nom_local(element.tag) == "si":
                chaines.append(texte_si(element))
                element.clear()
    return chaines


def texte_si(element: ET.Element) -> str:
    """Texte d'un <si> ou d'un <is> : <t> directs et <r><t>, sans <rPh>."""
    morceaux = []
    for enfant in element:
        local = nom_local(enfant.tag)
        if local == "t":
            morceaux.append(enfant.text or "")
        elif local == "r":
            morceaux.extend(t.text or "" for t in enfant if nom_local(t.tag) == "t")
    return "".join(morceaux)


def styles_dates(archive: zipfile.ZipFile, nom: str | None, plafond: int) -> frozenset[int]:
    """Indices de cellXfs dont le format numérique est une date ou une heure."""
    if nom is None or nom not in archive.namelist():
        return frozenset()
    racine = lire_xml(archive, nom, plafond)
    formats = {}
    for element in racine.iter():
        if nom_local(element.tag) == "numFmt":
            formats[int(element.get("numFmtId", "0"))] = element.get("formatCode", "")
    indices = set()
    for bloc in racine:
        if nom_local(bloc.tag) == "cellXfs":
            for position, xf in enumerate(x for x in bloc if nom_local(x.tag) == "xf"):
                identifiant = int(xf.get("numFmtId", "0"))
                if identifiant in FORMATS_DATE_INTEGRES or format_est_date(formats.get(identifiant, "")):
                    indices.add(position)
    return frozenset(indices)


def format_est_date(code: str) -> bool:
    """Un code de format personnalisé affiche-t-il une date ou une heure ?"""
    sans_texte = re.sub(r'"[^"]*"|\\.|_.|\*.', "", code)
    sans_crochets = re.sub(r"\[(?![hms]+\])[^\]]*\]", "", sans_texte, flags=re.IGNORECASE)
    section = sans_crochets.split(";", 1)[0]
    return bool(re.search(r"[dmyhs]", section, re.IGNORECASE)) and "General" not in section


@dataclass
class ContexteOoxml:
    """Ce qu'il faut savoir du classeur pour décoder ses cellules."""

    chaines: list[str]
    dates: frozenset[int]
    origine: dt.datetime


def lire_ooxml(chemin: Path, reglages: Reglages) -> dict[str, Any]:
    """Moteur stdlib pour .xlsx/.xlsm/.xltx/.xltm."""
    with ouvrir_archive(chemin) as archive:
        nom_classeur = partie_classeur(archive, reglages.max_octets)
        classeur = lire_xml(archive, nom_classeur, reglages.max_octets)
        rels = relations(archive, nom_classeur, reglages.max_octets)
        par_type = {type_rel: cible for type_rel, cible in rels.values()}
        contexte = ContexteOoxml(
            chaines=chaines_partagees(archive, par_type.get("sharedStrings"), reglages.max_octets),
            dates=styles_dates(archive, par_type.get("styles"), reglages.max_octets),
            origine=ORIGINE_1904 if est_1904(classeur) else ORIGINE_1900,
        )
        feuilles = [
            lire_feuille_ooxml(archive, declaration, rels, contexte, reglages)
            for declaration in declarations_feuilles(classeur)
        ]
    return {"moteur": "stdlib", "feuilles": feuilles}


def est_1904(classeur: ET.Element) -> bool:
    """Le classeur compte-t-il ses dates depuis 1904 ?"""
    for element in classeur:
        if nom_local(element.tag) == "workbookPr":
            return element.get("date1904", "0").lower() in ("1", "true")
    return False


def declarations_feuilles(classeur: ET.Element) -> list[dict[str, str]]:
    """Nom, état et identifiant de relation de chaque feuille, dans l'ordre."""
    declarations = []
    for element in classeur.iter():
        if nom_local(element.tag) == "sheet":
            declarations.append({
                "nom": element.get("name", "?"),
                "etat": element.get("state", "visible"),
                "rid": attribut_relation(element) or "",
            })
    return declarations


def lire_feuille_ooxml(
    archive: zipfile.ZipFile,
    declaration: dict[str, str],
    rels: dict[str, tuple[str, str]],
    contexte: ContexteOoxml,
    reglages: Reglages,
) -> dict[str, Any]:
    """Rapport d'une feuille ; les feuilles graphiques n'ont pas de cellules."""
    bilan = Bilan(declaration["nom"], declaration["etat"], reglages.max_cellules, reglages.apercu)
    type_rel, cible = rels.get(declaration["rid"], ("", ""))
    if type_rel != "worksheet":
        bilan.notes.append(f"feuille de type « {type_rel or 'inconnu'} » : pas de grille de cellules")
        return bilan.rapport()
    if cible not in archive.namelist():
        bilan.notes.append(f"partie {cible} absente de l'archive")
        return bilan.rapport()
    with ouvrir_partie(archive, cible, reglages.max_octets) as flux:
        for cellule in cellules_ooxml(flux, contexte):
            if not bilan.ajouter(cellule):
                break
    return bilan.rapport()


def cellules_ooxml(flux: Any, contexte: ContexteOoxml) -> Iterator[Cellule]:
    """Parcourt <sheetData> en flux, sans garder les lignes déjà lues."""
    ligne, colonne = 0, 0
    donnees: ET.Element | None = None
    for evenement, element in ET.iterparse(flux, events=("start", "end")):
        local = nom_local(element.tag)
        if evenement == "start":
            if local == "sheetData":
                donnees = element
            elif local == "row":
                ligne = int(element.get("r") or ligne + 1)
                colonne = 0
            continue
        if local == "c":
            ligne, colonne = position_cellule(element.get("r"), ligne, colonne)
            cellule = decoder_cellule(element, ligne, colonne, contexte)
            element.clear()
            if cellule is not None:
                yield cellule
        elif local == "row" and donnees is not None:
            donnees.clear()


def position_cellule(ref: str | None, ligne: int, colonne: int) -> tuple[int, int]:
    """Position d'après l'attribut r, ou la cellule suivante s'il manque."""
    if ref:
        correspondance = MOTIF_REFERENCE.fullmatch(ref)
        if correspondance:
            return int(correspondance.group(2)), numero_colonne(correspondance.group(1))
    return ligne, colonne + 1


def decoder_cellule(element: ET.Element, ligne: int, colonne: int, contexte: ContexteOoxml) -> Cellule | None:
    """Type, valeur et formule d'un élément <c> ; None s'il est vide."""
    brut, formule, en_ligne = None, None, None
    for enfant in element:
        local = nom_local(enfant.tag)
        if local == "v":
            brut = enfant.text
        elif local == "f":
            formule = texte_formule(enfant)
        elif local == "is":
            en_ligne = texte_si(enfant)
    type_cellule = element.get("t", "n")
    if type_cellule == "inlineStr":
        brut = en_ligne
    if brut is None or brut == "":
        if formule is None:
            return None
        return Cellule(ligne, colonne, "sans_valeur", None, formule)
    genre, valeur = interpreter_valeur(type_cellule, brut, element.get("s"), contexte)
    return Cellule(ligne, colonne, genre, valeur, formule)


def texte_formule(element: ET.Element) -> str:
    """Texte d'une formule ; une formule partagée sans texte est signalée."""
    if element.text:
        return "=" + element.text
    if element.get("t") == "shared":
        return f"(formule partagée si={element.get('si', '?')})"
    return "(formule vide)"


def interpreter_valeur(type_cellule: str, brut: str, style: str | None, contexte: ContexteOoxml) -> tuple[str, Any]:
    """Applique la sémantique de l'attribut t (et du style pour les dates)."""
    if type_cellule == "s":
        index = int(brut)
        texte = contexte.chaines[index] if 0 <= index < len(contexte.chaines) else ""
        return genre_texte(texte), texte
    if type_cellule in ("str", "inlineStr"):
        return genre_texte(brut), brut
    if type_cellule == "b":
        return "booleen", brut.strip() in ("1", "true")
    if type_cellule == "e":
        return "erreur", brut
    if type_cellule == "d":
        return "date", brut
    nombre = float(brut)
    if style is not None and int(style) in contexte.dates:
        date = date_depuis_serie(nombre, contexte.origine)
        if date is not None:
            return "date", date
    return "nombre", nombre


# --------------------------------------------------------------------------- #
# OpenDocument (stdlib)

def lire_odf(chemin: Path, reglages: Reglages) -> dict[str, Any]:
    """Moteur stdlib pour .ods/.ots : content.xml lu en flux."""
    with ouvrir_archive(chemin) as archive:
        if "content.xml" not in archive.namelist():
            raise ErreurClasseur(f"{chemin.name} : content.xml absent, ce n'est pas un document OpenDocument")
        with ouvrir_partie(archive, "content.xml", reglages.max_octets) as flux:
            feuilles = feuilles_odf(flux, reglages)
    return {"moteur": "stdlib", "feuilles": feuilles}


def attribut_odf(element: ET.Element, espace: str, nom: str) -> str | None:
    """Attribut qualifié par l'URI de son espace de noms."""
    return element.get(f"{{{espace}}}{nom}")


@dataclass
class EtatOdf:
    """Position courante pendant le parcours de content.xml."""

    masques: set[str] = field(default_factory=set)
    style_courant: str = ""
    bilan: Bilan | None = None
    ligne: int = 0
    saturee: bool = False
    pile: list[ET.Element] = field(default_factory=list)


def feuilles_odf(flux: Any, reglages: Reglages) -> list[dict[str, Any]]:
    """Parcourt les tables ; les répétitions vides ne sont jamais dépliées."""
    etat = EtatOdf()
    rapports: list[dict[str, Any]] = []
    for evenement, element in ET.iterparse(flux, events=("start", "end")):
        if evenement == "start":
            etat.pile.append(element)
            debuter_element_odf(etat, element, reglages)
            continue
        etat.pile.pop()
        if element.tag == f"{{{NS_STYLE}}}table-properties" and attribut_odf(element, NS_TABLE, "display") == "false":
            etat.masques.add(etat.style_courant)
        elif element.tag == f"{{{NS_TABLE}}}table-row" and etat.bilan is not None:
            traiter_ligne_odf(etat, element)
            liberer(etat.pile, element)
        elif element.tag == f"{{{NS_TABLE}}}table" and etat.bilan is not None:
            rapports.append(etat.bilan.rapport())
            etat.bilan = None
            liberer(etat.pile, element)
    return rapports


def liberer(pile: list[ET.Element], element: ET.Element) -> None:
    """Détache un élément déjà traité de son parent pour libérer la mémoire."""
    element.clear()
    if pile:
        pile[-1].remove(element)


def debuter_element_odf(etat: EtatOdf, element: ET.Element, reglages: Reglages) -> None:
    """Réagit à l'ouverture d'un style ou d'une table."""
    if element.tag == f"{{{NS_STYLE}}}style":
        etat.style_courant = attribut_odf(element, NS_STYLE, "name") or ""
    elif element.tag == f"{{{NS_TABLE}}}table" and etat.bilan is None:
        style = attribut_odf(element, NS_TABLE, "style-name") or ""
        etat.bilan = Bilan(
            attribut_odf(element, NS_TABLE, "name") or "?",
            "hidden" if style in etat.masques else "visible",
            reglages.max_cellules,
            reglages.apercu,
        )
        etat.ligne = 0
        etat.saturee = False


def traiter_ligne_odf(etat: EtatOdf, rangee: ET.Element) -> None:
    """Émet les cellules non vides d'une ligne, répétitions comprises (bornées)."""
    repetitions = int(attribut_odf(rangee, NS_TABLE, "number-rows-repeated") or "1")
    cellules = cellules_ligne_odf(rangee)
    premiere = etat.ligne + 1
    etat.ligne += repetitions
    if not cellules or etat.saturee or etat.bilan is None:
        return
    for decalage in range(repetitions):
        for colonne, genre, valeur, formule in cellules:
            if not etat.bilan.ajouter(Cellule(premiere + decalage, colonne, genre, valeur, formule)):
                etat.saturee = True
                return


def cellules_ligne_odf(rangee: ET.Element) -> list[tuple[int, str, Any, str | None]]:
    """(colonne, genre, valeur, formule) des cellules non vides d'une ligne."""
    resultat = []
    colonne = 0
    for cellule in rangee:
        if cellule.tag not in (f"{{{NS_TABLE}}}table-cell", f"{{{NS_TABLE}}}covered-table-cell"):
            continue
        repetitions = int(attribut_odf(cellule, NS_TABLE, "number-columns-repeated") or "1")
        decodee = decoder_cellule_odf(cellule)
        if decodee is not None:
            resultat.extend((colonne + 1 + k, *decodee) for k in range(min(repetitions, 16384)))
        colonne += repetitions
    return resultat


def decoder_cellule_odf(cellule: ET.Element) -> tuple[str, Any, str | None] | None:
    """Genre, valeur et formule d'une cellule OpenDocument ; None si vide."""
    formule = attribut_odf(cellule, NS_TABLE, "formula")
    type_valeur = attribut_odf(cellule, NS_OFFICE, "value-type")
    texte = texte_cellule_odf(cellule)
    if attribut_odf(cellule, NS_CALCEXT, "value-type") == "error":
        return "erreur", texte, formule
    if type_valeur is None:
        if texte:
            return genre_texte(texte), texte, formule
        return ("sans_valeur", None, formule) if formule else None
    return genre_odf(type_valeur, cellule, texte, formule)


def genre_odf(type_valeur: str, cellule: ET.Element, texte: str, formule: str | None) -> tuple[str, Any, str | None]:
    """Traduit office:value-type en genre commun."""
    if type_valeur in ("float", "percentage", "currency"):
        return "nombre", float(attribut_odf(cellule, NS_OFFICE, "value") or "nan"), formule
    if type_valeur == "date":
        return "date", attribut_odf(cellule, NS_OFFICE, "date-value") or texte, formule
    if type_valeur == "time":
        return "date", attribut_odf(cellule, NS_OFFICE, "time-value") or texte, formule
    if type_valeur == "boolean":
        return "booleen", attribut_odf(cellule, NS_OFFICE, "boolean-value") == "true", formule
    if formule and est_code_erreur(texte):
        return "erreur", texte, formule
    return genre_texte(texte), texte, formule


def texte_cellule_odf(cellule: ET.Element) -> str:
    """Texte visible des paragraphes d'une cellule, annotations exclues."""
    paragraphes = [texte_paragraphe_odf(p) for p in cellule if p.tag in (f"{{{NS_TEXT}}}p", f"{{{NS_TEXT}}}h")]
    return "\n".join(paragraphes)


def texte_paragraphe_odf(element: ET.Element) -> str:
    """Texte d'un text:p avec text:s, text:tab et text:line-break."""
    morceaux = [element.text or ""]
    for enfant in element:
        local = nom_local(enfant.tag)
        if local == "s":
            morceaux.append(" " * int(attribut_odf(enfant, NS_TEXT, "c") or "1"))
        elif local == "tab":
            morceaux.append("\t")
        elif local == "line-break":
            morceaux.append("\n")
        elif local != "annotation":
            morceaux.append(texte_paragraphe_odf(enfant))
        morceaux.append(enfant.tail or "")
    return "".join(morceaux)


# --------------------------------------------------------------------------- #
# Texte délimité (csv/tsv)

def lire_delimite(chemin: Path, reglages: Reglages) -> dict[str, Any]:
    """Lit un .csv/.tsv en flux après détection d'encodage et de délimiteur."""
    with chemin.open("rb") as brut:
        echantillon = brut.read(65536)
    encodage = deviner_encodage(echantillon, chemin.name)
    with chemin.open("r", encoding=encodage, errors="replace", newline="") as texte:
        return lire_texte_delimite(texte, encodage, chemin.suffix.lower(), reglages)


def lire_texte_delimite(texte: Any, encodage: str, suffixe: str, reglages: Reglages) -> dict[str, Any]:
    """Analyse un flux texte délimité déjà ouvert (relu depuis le début)."""
    debut = texte.read(65536)
    texte.seek(0)
    delimiteur, source = choisir_delimiteur(debut, suffixe, reglages.delimiteur)
    bilan = Bilan("(texte)", "visible", reglages.max_cellules, reglages.apercu)
    lignes = lignes_bornees(texte, reglages.max_octets, bilan)
    irregulieres = remplir_bilan_csv(bilan, csv.reader(lignes, delimiter=delimiteur), delimiteur)
    if bilan.nb_formules:
        bilan.notes.append("en csv, une cellule qui commence par « = » est du texte qu'un tableur "
                           "exécutera à l'ouverture (injection de formule)")
    remplacements = sum(str(v).count("\ufffd") for ligne in bilan.apercu.values() for v in ligne.values())
    details = {"encodage": encodage, "delimiteur": delimiteur, "delimiteur_source": source,
               "lignes_irregulieres": irregulieres, "caracteres_remplaces_apercu": remplacements}
    return {"moteur": "stdlib", "feuilles": [bilan.rapport()], "texte_delimite": details}


def lignes_bornees(texte: Any, plafond: int, bilan: Bilan) -> Iterator[str]:
    """Lignes du flux, chacune et toutes ensemble plafonnées."""
    lus = 0
    while True:
        ligne = texte.readline(LIGNE_MAX)
        if not ligne:
            return
        lus += len(ligne)
        if lus > plafond or (len(ligne) == LIGNE_MAX and not ligne.endswith(("\n", "\r"))):
            bilan.tronque = True
            bilan.notes.append(f"lecture arrêtée après {lus} caractères (ligne de plus de "
                               f"{LIGNE_MAX} caractères ou plafond --max-octets)")
            return
        yield ligne


def deviner_encodage(echantillon: bytes, nom: str) -> str:
    """utf-8 (BOM ou non) sinon cp1252 ; refuse un fichier binaire."""
    if b"\x00" in echantillon:
        raise ErreurClasseur(f"{nom} contient des octets nuls : fichier binaire, pas du texte délimité")
    if echantillon.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    try:
        echantillon.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError as exc:
        if exc.start >= len(echantillon) - 3:
            return "utf-8"
    return "cp1252"


def choisir_delimiteur(debut: str, suffixe: str, impose: str | None) -> tuple[str, str]:
    """Délimiteur imposé, sinon deviné par csv.Sniffer, sinon selon l'extension."""
    if impose:
        return impose, "impose"
    try:
        return csv.Sniffer().sniff(debut, delimiters=CANDIDATS_DELIMITEUR).delimiter, "csv.Sniffer"
    except csv.Error:
        pass
    frequent = delimiteur_frequent(debut)
    if frequent is not None:
        return frequent, "fréquence par ligne (csv.Sniffer indécis)"
    return ("\t" if suffixe == ".tsv" else ","), "extension (csv.Sniffer indécis)"


def delimiteur_frequent(debut: str) -> str | None:
    """Candidat présent avec le même compte sur le plus de lignes (50 premières)."""
    lignes = [ligne for ligne in debut.splitlines()[:50] if ligne.strip()]
    meilleur, score_max = None, 0
    for candidat in CANDIDATS_DELIMITEUR:
        comptes = Counter(ligne.count(candidat) for ligne in lignes)
        comptes.pop(0, None)
        score = max(comptes.values(), default=0)
        if score > score_max:
            meilleur, score_max = candidat, score
    return meilleur


def remplir_bilan_csv(bilan: Bilan, lecteur: Iterator[list[str]], delimiteur: str) -> dict[str, Any]:
    """Alimente le bilan ; rend le décompte des lignes de largeur irrégulière."""
    largeur, irregulieres, exemples = None, 0, []
    for numero, champs in enumerate(lecteur, start=1):
        if not any(champ.strip() for champ in champs):
            continue
        largeur = len(champs) if largeur is None else largeur
        if len(champs) != largeur:
            irregulieres += 1
            if len(exemples) < LIMITE_EXEMPLES:
                exemples.append({"ligne": numero, "champs": len(champs), "attendus": largeur})
        if not all(bilan.ajouter(c) for c in cellules_csv(numero, champs, delimiteur)):
            break
    return {"nombre": irregulieres, "exemples": exemples}


def cellules_csv(numero: int, champs: list[str], delimiteur: str) -> Iterator[Cellule]:
    """Cellules non vides d'une ligne délimitée, typées par leur texte."""
    for colonne, champ in enumerate(champs, start=1):
        if champ.strip():
            genre, valeur, formule = genre_csv(champ, delimiteur)
            yield Cellule(numero, colonne, genre, valeur, formule)


def genre_csv(champ: str, delimiteur: str) -> tuple[str, Any, str | None]:
    """Infère le type d'un champ textuel."""
    propre = champ.strip()
    if est_code_erreur(propre):
        return "erreur", propre, None
    if propre.startswith("="):
        return "texte", propre, propre
    if propre.lower() in BOOLEENS_TEXTE:
        return "booleen", propre, None
    if MOTIF_NOMBRE.fullmatch(propre) and ("," not in propre or delimiteur != ","):
        return "nombre", float(propre.replace(",", ".")), None
    if MOTIF_DATE_TEXTE.fullmatch(propre):
        return "date", propre, None
    return "texte", champ, None


# --------------------------------------------------------------------------- #
# Moteurs optionnels

def lire_openpyxl(chemin: Path, reglages: Reglages) -> dict[str, Any]:
    """Moteur openpyxl : deux lectures seules (formules, valeurs en cache)."""
    with ouvrir_archive(chemin) as archive:
        verifier_tailles(archive, reglages.max_octets)
    formules = openpyxl.load_workbook(chemin, read_only=True, data_only=False)
    valeurs = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    try:
        feuilles = [feuille_openpyxl(formules[nom], valeurs[nom], reglages) for nom in formules.sheetnames]
    finally:
        formules.close()
        valeurs.close()
    return {"moteur": "openpyxl", "feuilles": feuilles}


def feuille_openpyxl(feuille_f: Any, feuille_v: Any, reglages: Reglages) -> dict[str, Any]:
    """Rapport d'une feuille openpyxl ; la dimension déclarée est ignorée."""
    bilan = Bilan(feuille_f.title, feuille_f.sheet_state, reglages.max_cellules, reglages.apercu)
    if not hasattr(feuille_f, "reset_dimensions"):
        bilan.notes.append("feuille de type « chartsheet » : pas de grille de cellules")
        return bilan.rapport()
    feuille_f.reset_dimensions()
    feuille_v.reset_dimensions()
    for rangee_f, rangee_v in zip(feuille_f.iter_rows(), feuille_v.iter_rows()):
        for cellule_f, cellule_v in zip(rangee_f, rangee_v):
            cellule = cellule_openpyxl(cellule_f, cellule_v)
            if cellule is not None and not bilan.ajouter(cellule):
                return bilan.rapport()
    return bilan.rapport()


def cellule_openpyxl(cellule_f: Any, cellule_v: Any) -> Cellule | None:
    """Fusionne la vue formules et la vue valeurs d'une même cellule."""
    if cellule_f.value is None or not hasattr(cellule_f, "row"):
        return None
    formule = cellule_f.value if cellule_f.data_type == "f" else None
    valeur = cellule_v.value
    if valeur is None:
        return Cellule(cellule_f.row, cellule_f.column, "sans_valeur", None, formule)
    return Cellule(cellule_f.row, cellule_f.column, genre_python(valeur, cellule_v.data_type), valeur_python(valeur), formule)


def genre_python(valeur: Any, type_donnee: str | None) -> str:
    """Genre commun d'une valeur Python rendue par openpyxl ou python-calamine."""
    if type_donnee == "e":
        return "erreur"
    if isinstance(valeur, bool):
        return "booleen"
    if isinstance(valeur, (int, float)):
        return "nombre"
    if isinstance(valeur, (dt.date, dt.time, dt.timedelta)):
        return "date"
    return genre_texte(str(valeur))


def valeur_python(valeur: Any) -> Any:
    """Dates en ISO, le reste tel quel."""
    if isinstance(valeur, dt.datetime):
        return valeur.isoformat(timespec="seconds") if valeur.time() != dt.time(0, 0) else valeur.date().isoformat()
    if isinstance(valeur, (dt.date, dt.time)):
        return valeur.isoformat()
    if isinstance(valeur, dt.timedelta):
        return str(valeur)
    return valeur


def lire_calamine(chemin: Path, reglages: Reglages) -> dict[str, Any]:
    """Moteur python-calamine : valeurs seules, ni formules ni erreurs."""
    classeur = python_calamine.CalamineWorkbook.from_path(str(chemin))
    try:
        feuilles = [feuille_calamine(classeur, meta, reglages) for meta in classeur.sheets_metadata]
    finally:
        classeur.close()
    return {"moteur": "python-calamine", "feuilles": feuilles}


def feuille_calamine(classeur: Any, meta: Any, reglages: Reglages) -> dict[str, Any]:
    """Rapport d'une feuille calamine (formules et erreurs inconnues)."""
    visibilite = str(meta.visible).rsplit(".", 1)[-1]
    etat = {"Visible": "visible", "Hidden": "hidden", "VeryHidden": "veryHidden"}.get(visibilite, visibilite)
    bilan = Bilan(meta.name, etat, reglages.max_cellules, reglages.apercu,
                  formules_connues=False, erreurs_connues=False)
    bilan.notes.append("python-calamine lit les cellules en erreur comme vides et ne donne pas les formules")
    if not str(meta.typ).endswith("WorkSheet"):
        bilan.notes.append(f"feuille de type « {str(meta.typ).rsplit('.', 1)[-1]} » : pas de grille de cellules")
        return bilan.rapport()
    feuille = classeur.get_sheet_by_name(meta.name)
    if feuille.start is None:
        return bilan.rapport()
    ligne0, colonne0 = feuille.start
    for decalage, rangee in enumerate(feuille.iter_rows()):
        for position, valeur in enumerate(rangee):
            if valeur == "" or valeur is None:
                continue
            cellule = Cellule(ligne0 + decalage + 1, colonne0 + position + 1,
                              genre_python(valeur, None), valeur_python(valeur), None)
            if not bilan.ajouter(cellule):
                return bilan.rapport()
    return bilan.rapport()


# --------------------------------------------------------------------------- #
# Aiguillage

def choisir_lecteur(suffixe: str, moteur: str) -> tuple[Callable[[Path, Reglages], dict[str, Any]], str | None]:
    """Fonction de lecture pour une extension ; second élément : repli annoncé."""
    if moteur == "calamine" and suffixe in EXTENSIONS_OOXML + EXTENSIONS_ODF + EXTENSIONS_BINAIRES:
        if python_calamine is None:
            raise ErreurClasseur("--moteur calamine demandé mais python-calamine n'est pas installé")
        return lire_calamine, None
    if suffixe in EXTENSIONS_BINAIRES:
        if python_calamine is None:
            raise ErreurClasseur(f"format binaire {suffixe} : illisible sans python-calamine (absent)")
        return lire_calamine, None
    if suffixe in EXTENSIONS_OOXML:
        if moteur == "openpyxl" and openpyxl is None:
            raise ErreurClasseur("--moteur openpyxl demandé mais openpyxl n'est pas installé")
        if moteur in ("auto", "openpyxl") and openpyxl is not None:
            return lire_openpyxl, None
        return lire_ooxml, ("openpyxl absent" if moteur == "auto" else None)
    if suffixe in EXTENSIONS_ODF:
        return lire_odf, None
    if suffixe in EXTENSIONS_TEXTE:
        return lire_delimite, None
    raise ErreurClasseur(f"extension « {suffixe or '(aucune)'} » non prise en charge")


def examiner_classeur(chemin: Path, reglages: Reglages) -> tuple[dict[str, Any], str | None]:
    """Rapport d'un classeur ; un échec d'openpyxl se replie sur la stdlib, en le disant."""
    lecteur, repli = choisir_lecteur(chemin.suffix.lower(), reglages.moteur)
    if lecteur is lire_calamine:
        resultat = lire_calamine_garde(chemin, reglages)
    elif lecteur is lire_openpyxl:
        resultat = lire_openpyxl_garde(chemin, reglages)
    else:
        resultat = lecteur(chemin, reglages)
    resultat["format"] = chemin.suffix.lower().lstrip(".")
    if chemin.suffix.lower() in EXTENSIONS_OOXML:
        resultat.update(inventaire_ooxml(chemin))
    return resultat, repli


def lire_openpyxl_garde(chemin: Path, reglages: Reglages) -> dict[str, Any]:
    """openpyxl, avec repli stdlib annoncé si la bibliothèque échoue."""
    try:
        return lire_openpyxl(chemin, reglages)
    except ErreurClasseur:
        raise
    except ERREURS_LECTURE + (TypeError, AttributeError, IndexError) as exc:
        print(f"openpyxl a échoué sur {chemin.name} ({type(exc).__name__}: {exc}) ; repli stdlib", file=sys.stderr)
        resultat = lire_ooxml(chemin, reglages)
        resultat["repli"] = f"openpyxl : {type(exc).__name__}"
        return resultat


def lire_calamine_garde(chemin: Path, reglages: Reglages) -> dict[str, Any]:
    """python-calamine, ses erreurs (et ses paniques Rust) traduites en refus."""
    try:
        return lire_calamine(chemin, reglages)
    except python_calamine.CalamineError as exc:
        raise ErreurClasseur(f"{chemin.name} : python-calamine refuse ({exc})") from exc
    except BaseException as exc:
        if type(exc).__name__ != "PanicException":
            raise
        raise ErreurClasseur(f"{chemin.name} : python-calamine a paniqué ({exc})") from exc


def lister_classeurs(dossier: Path, plafond: int) -> tuple[list[Path], list[str]]:
    """Classeurs d'un dossier (récursif, trié) ; fichiers verrous ~$ écartés."""
    acceptees = EXTENSIONS_OOXML + EXTENSIONS_ODF + EXTENSIONS_TEXTE + EXTENSIONS_BINAIRES
    retenus, ecartes = [], []
    for chemin in sorted(dossier.rglob("*")):
        if not chemin.is_file() or chemin.suffix.lower() not in acceptees:
            continue
        if chemin.name.startswith("~$") or chemin.name.startswith(".~lock"):
            ecartes.append(f"{chemin.relative_to(dossier)} (fichier verrou)")
        elif len(retenus) < plafond:
            retenus.append(chemin)
        else:
            ecartes.append(f"{chemin.relative_to(dossier)} (au-delà de --max-fichiers)")
    return retenus, ecartes


def examiner_chemin(cible: Path, reglages: Reglages, max_fichiers: int) -> dict[str, Any]:
    """Examine un fichier ou tous les classeurs d'un dossier."""
    if not cible.exists():
        raise ErreurClasseur(f"chemin introuvable : {cible}")
    if cible.is_dir():
        fichiers, ecartes = lister_classeurs(cible, max_fichiers)
        base = cible
    else:
        fichiers, ecartes, base = [cible], [], cible.parent
    classeurs, illisibles, replis = [], [], set()
    for fichier in fichiers:
        nom = str(fichier.relative_to(base))
        try:
            rapport, repli = examiner_classeur(fichier, reglages)
        except ERREURS_LECTURE as exc:
            if not cible.is_dir():
                raise ErreurClasseur(f"{nom} : {exc}", getattr(exc, "code", 2)) from exc
            illisibles.append({"fichier": nom, "erreur": f"{type(exc).__name__}: {exc}"})
            continue
        if repli:
            replis.add(repli)
        classeurs.append({"fichier": nom, **rapport})
    return {"classeurs": classeurs, "illisibles": illisibles, "ecartes": ecartes, "replis": sorted(replis)}


def examiner_texte(texte: str, reglages: Reglages) -> dict[str, Any]:
    """Analyse un texte délimité passé en ligne de commande."""
    rapport = lire_texte_delimite(io.StringIO(texte), "utf-8", ".csv", reglages)
    rapport["format"] = "csv"
    return {"classeurs": [{"fichier": "(--texte)", **rapport}], "illisibles": [], "ecartes": [], "replis": []}


def assembler(resultat: dict[str, Any]) -> dict[str, Any]:
    """Objet JSON final ; `denominateur` en tête (classeurs réellement examinés)."""
    noms = [c["fichier"] for c in resultat["classeurs"]] + [i["fichier"] for i in resultat["illisibles"]]
    feuilles = [f for c in resultat["classeurs"] for f in c["feuilles"]]
    moteurs = sorted({c["moteur"] for c in resultat["classeurs"]}) or ["stdlib"]
    return {
        "denominateur": len(noms),
        "examines": noms[:LIMITE_EXAMINES],
        "examines_tronques": len(noms) > LIMITE_EXAMINES,
        "moteur": "+".join(moteurs),
        "feuilles_examinees": len(feuilles),
        "erreurs_detectables": all(f["nb_erreurs"] is not None for f in feuilles),
        "cellules_en_erreur": sum(f["nb_erreurs"] or 0 for f in feuilles),
        "formules": sum(f["nb_formules"] or 0 for f in feuilles),
        "illisibles": resultat["illisibles"],
        "ecartes": resultat["ecartes"],
        "classeurs": resultat["classeurs"],
        "contrat": extraire_contrat(__doc__ or ""),
    }


def code_de_sortie(sortie: dict[str, Any]) -> int:
    """1 si une cellule est en erreur ou un classeur illisible, sinon 0."""
    return 1 if sortie["cellules_en_erreur"] or sortie["illisibles"] else 0


# --------------------------------------------------------------------------- #
# Présentation

def afficher_json(sortie: dict[str, Any]) -> None:
    """Un seul objet JSON sur stdout."""
    print(json.dumps(sortie, ensure_ascii=False, indent=2, default=str))


def afficher_humain(sortie: dict[str, Any]) -> None:
    """Résumé lisible, une feuille par bloc."""
    print(f"{sortie['denominateur']} classeur(s) examiné(s), moteur {sortie['moteur']} ; "
          f"{sortie['cellules_en_erreur']} cellule(s) en erreur, {sortie['formules']} formule(s)")
    for classeur in sortie["classeurs"]:
        print(f"\n{classeur['fichier']} ({classeur['format']}, moteur {classeur['moteur']})")
        for feuille in classeur["feuilles"]:
            afficher_feuille(feuille)
    for illisible in sortie["illisibles"]:
        print(f"\nILLISIBLE {illisible['fichier']} : {illisible['erreur']}")


def afficher_feuille(feuille: dict[str, Any]) -> None:
    """Bloc lisible d'une feuille."""
    etat = "" if feuille["etat"] == "visible" else f" [{feuille['etat']}]"
    print(f"  feuille « {feuille['nom']} »{etat} : {feuille['dimension'] or 'vide'}, "
          f"{feuille['cellules_non_vides']} cellule(s) non vide(s)")
    if feuille["entetes"]:
        print(f"    en-têtes (ligne {feuille['ligne_entete']}) : {' | '.join(feuille['entetes'])}")
    for colonne in feuille["colonnes_detail"]:
        types = ", ".join(f"{g} {n}" for g, n in colonne["types"].items()) or "aucune donnée"
        alerte = "  (mixte)" if colonne["mixte"] else ""
        print(f"    {colonne['colonne']} {colonne['entete'][:30]!r} : {types}, vides {colonne['vides']}{alerte}")
    for erreur in feuille["erreurs"] or []:
        print(f"    ERREUR {erreur['cellule']} = {erreur['valeur']}  {erreur['formule'] or ''}")
    if feuille["nb_formules"]:
        print(f"    {feuille['nb_formules']} formule(s), dont {feuille['formules_sans_valeur']} sans valeur calculée")
    for note in feuille["notes"]:
        print(f"    note : {note}")


# --------------------------------------------------------------------------- #
# Interface

def construire_analyseur() -> argparse.ArgumentParser:
    """Analyseur d'arguments en français."""
    analyseur = argparse.ArgumentParser(
        description="Décrit un classeur (.xlsx, .xlsm, .ods, .csv, .tsv ; .xls/.xlsb avec "
                    "python-calamine) : feuilles, dimensions, en-têtes, types par colonne, "
                    "cellules en erreur, formules. Code 1 si une cellule est en erreur.",
        epilog="Exemple : python3 lire_tableur.py rapports/ventes.xlsx --json",
    )
    analyseur.add_argument("chemin", nargs="?", help="classeur, ou dossier dont examiner les classeurs")
    analyseur.add_argument("--texte", help="contenu csv passé directement (au lieu d'un chemin)")
    analyseur.add_argument("--json", action="store_true", help="un seul objet JSON sur stdout")
    analyseur.add_argument("--racine", help="dossier de base des chemins relatifs (défaut : dossier courant)")
    analyseur.add_argument("--moteur", choices=("auto", "stdlib", "openpyxl", "calamine"), default="auto",
                           help="auto : openpyxl pour .xlsx s'il est installé, stdlib sinon")
    analyseur.add_argument("--delimiteur", help="délimiteur csv imposé (sinon deviné)")
    analyseur.add_argument("--apercu", type=int, default=3, help="lignes de données montrées par feuille (défaut 3)")
    analyseur.add_argument("--max-cellules", type=int, default=2_000_000, help="plafond de cellules lues par feuille")
    analyseur.add_argument("--max-octets", type=int, default=256 * 1024 * 1024,
                           help="plafond d'octets décompressés par partie (garde contre les bombes zip)")
    analyseur.add_argument("--max-fichiers", type=int, default=500, help="plafond de classeurs examinés dans un dossier")
    return analyseur


def resoudre(chemin: str, racine: str | None) -> Path:
    """Chemin relatif résolu contre --racine (ou le dossier courant)."""
    brut = Path(chemin)
    if brut.is_absolute() or racine is None:
        return brut
    return Path(racine) / brut


def annoncer_repli(moteur: str, replis: list[str]) -> None:
    """Une ligne sur stderr quand une bibliothèque optionnelle manque."""
    if "openpyxl absent" in replis and moteur == "auto":
        print("openpyxl absent : lecture .xlsx en mode dégradé stdlib (zipfile + xml.etree)", file=sys.stderr)


def main() -> int:
    """Point d'entrée : 0 rien à signaler, 1 erreur trouvée, 2 entrée invalide, 3 rien à examiner."""
    args = construire_analyseur().parse_args()
    if (args.chemin is None) == (args.texte is None):
        print("donner soit un chemin, soit --texte (et pas les deux)", file=sys.stderr)
        return 2
    if args.delimiteur is not None and len(args.delimiteur) != 1:
        print("--delimiteur attend un seul caractère", file=sys.stderr)
        return 2
    reglages = Reglages(args.moteur, max(args.max_cellules, 1), max(args.max_octets, 1024),
                        max(args.apercu, 0), args.delimiteur)
    try:
        if args.texte is not None:
            resultat = examiner_texte(args.texte, reglages)
        else:
            resultat = examiner_chemin(resoudre(args.chemin, args.racine), reglages, max(args.max_fichiers, 1))
    except ErreurClasseur as exc:
        print(f"lire_tableur : {exc}", file=sys.stderr)
        if args.json:
            afficher_json({"denominateur": 0, "examines": [], "moteur": "stdlib", "erreur": str(exc), "code": exc.code})
        return exc.code
    annoncer_repli(args.moteur, resultat["replis"])
    sortie = assembler(resultat)
    if not sortie["erreurs_detectables"]:
        print("python-calamine lit les cellules en erreur comme vides : un code 0 ne prouve pas "
              "l'absence d'erreur dans ces feuilles", file=sys.stderr)
    if args.json:
        afficher_json(sortie)
    else:
        afficher_humain(sortie)
    if sortie["denominateur"] == 0:
        print("dénominateur nul : aucun classeur trouvé, rien à examiner", file=sys.stderr)
        return 3
    return code_de_sortie(sortie)


if __name__ == "__main__":
    raise SystemExit(main())
